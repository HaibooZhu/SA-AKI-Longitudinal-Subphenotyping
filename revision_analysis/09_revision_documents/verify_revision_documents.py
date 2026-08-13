#!/usr/bin/env python3
"""Fail-closed structural and wording checks for final JTIM revision documents."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageChops


REQUIRED_FILES = (
    "JTIM_Revised_Manuscript_clean.docx",
    "JTIM_Revised_Manuscript_highlight.docx",
    "JTIM_Revised_Supplementary_material.docx",
    "JTIM_Point-by-point_Response.docx",
    "JTIM_逐点回复_中文工作版.docx",
    "JTIM_Revised_Cover_Letter.docx",
    "JTIM_Revised_Highlights.docx",
    "JTIM_Revised_Title_Page.docx",
)
WORDING_FILES = (
    "JTIM_Revised_Manuscript_clean.docx",
    "JTIM_Revised_Manuscript_highlight.docx",
    "JTIM_Revised_Highlights.docx",
    "JTIM_Revised_Cover_Letter.docx",
    "JTIM_Point-by-point_Response.docx",
)
PROHIBITED_EXACT_CLAIMS = (
    "The 3-cluster solution consistently achieved the lowest deviance and acceptable convergence across all datasets.",
    "actionable bedside classifier",
    "supports clinical deployment",
    "precision fluid therapy",
)
CONTENT_FILES = (
    "JTIM_Revised_Manuscript_clean.docx",
    "JTIM_Revised_Manuscript_highlight.docx",
    "JTIM_Point-by-point_Response.docx",
)
SUPPLEMENT_FIGURE_SOURCES = {
    "figure_s2": (
        "Figure S2.",
        "W3_deep_k_stability/Figure_S2_cross_cohort_k_stability.png",
    ),
    "figure_s10": (
        "Figure S10.",
        "W3_cross_cohort_robustness/cross_cohort_robustness_matrix.png",
    ),
    "figure_s11a": (
        "Figure S11a.",
        "W3_cluster_robustness/documented_windows_cluster_sensitivity.png",
    ),
    "figure_s11b": (
        "Figure S11b.",
        "W3_cluster_robustness/high_coverage_cluster_sensitivity.png",
    ),
    "figure_s12": (
        "Figure S12.",
        "W4_classifier_validation/Figure_S12_archived_model_calibration_comparator.png",
    ),
    "figure_s13": (
        "Figure S13.",
        "W5_independent_outcomes/W5_adjusted_outcomes_forest.png",
    ),
    "figure_s14": (
        "Figure S14.",
        "W6_diuretic_exploratory/W6_early_diuretic_response_forest.png",
    ),
}
REQUIRED_CONTENT_PATTERNS: dict[str, tuple[str, ...]] = {
    "aumc_uses_three_clustering_variables": (
        r"(?:AmsterdamUMCdb|AUMC).{0,450}(?:other three (?:responses|variables)|three (?:renal )?(?:responses|variables))",
    ),
    "eicu_k3_initialization_sensitivity": (
        r"eICU(?:-CRD)?.{0,260}initialization sensitivity",
    ),
    "high_coverage_uo_does_not_preserve_pw": (
        r"(?:≥50%|50%-coverage).{0,500}(?:PW component collapsed|zero (?:median )?PW (?:prevalence|fraction)|PW (?:class|component).{0,60}collapse)",
    ),
    "classifier_has_no_external_incremental_advantage": (
        r"(?:no external incremental (?:value|advantage)|did not improve.{0,140}externally)",
    ),
    "diuretic_analysis_is_exploratory": (
        r"diuretic.{0,450}(?:secondary.{0,80}exploratory|exploratory)",
    ),
    "complete_followup_and_missing_data_cautions": (
        r"complete(?: 30-window)? follow-up.{0,280}(?:caution|sensitivity)",
        r"(?:missing-data.{0,180}(?:sensitivity|rule|caution|refit)|sensitivity.{0,120}missing-data)",
    ),
    "k3_is_not_unanimous_or_uniquely_optimal": (
        r"(?:did not (?:select|provide unanimous support for) k=3|not select k=3 unanimously|k=3.{0,100}not.{0,100}(?:unique|unanimous)|does not establish k=3 as (?:a )?unique)",
    ),
}
NEGATION_GUARDS = re.compile(
    r"\b(?:remove(?:d)?|not|no longer|without|unsupported|cannot|does not)\b",
    flags=re.IGNORECASE,
)


def parse_args() -> argparse.Namespace:
    analysis = Path(__file__).resolve().parents[2]
    project = analysis.parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--package-dir",
        type=Path,
        default=(
            project
            / "01_manuscript/14、Journal of Translational Internal Medicine/revision_package_20260805"
        ),
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=analysis / "02_revision_outputs/reports/W8_revision_document_qa",
    )
    parser.add_argument(
        "--fresh-k-summary",
        type=Path,
        default=(
            analysis
            / "02_revision_outputs/reports/W3_fresh_k_grid/fresh_k_grid_summary.csv"
        ),
    )
    parser.add_argument(
        "--figure-s2-source",
        type=Path,
        default=(
            analysis
            / "02_revision_outputs/reports/W3_deep_k_stability/Figure_S2_cross_cohort_k_stability.png"
        ),
    )
    parser.add_argument(
        "--report-root",
        type=Path,
        default=analysis / "02_revision_outputs/reports",
    )
    return parser.parse_args()


def default_figure_sources(report_root: Path) -> dict[str, Path]:
    return {
        figure_id: report_root / relative_path
        for figure_id, (_, relative_path) in SUPPLEMENT_FIGURE_SOURCES.items()
    }


def document_text(path: Path) -> str:
    document = Document(path)
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        blocks.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(blocks)


def author_response_text(path: Path) -> str:
    """Return author-authored response paragraphs, excluding quoted comments and locations."""
    document = Document(path)
    collected: list[str] = []
    in_response = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text == "Response":
            in_response = True
            continue
        if (
            text in {
                "Revised text in the manuscript",
                "Changes in the manuscript and supporting evidence",
            }
            or paragraph.style.name.startswith("Heading")
        ):
            in_response = False
        elif in_response and text:
            collected.append(text)
    return "\n".join(collected)


def prohibited_claim_hits(text: str) -> list[str]:
    """Find legacy claims while ignoring explicit statements that they were removed."""
    hits: list[str] = []
    sentences = re.split(r"(?<=[.!?])\s+|\n+", text)
    for claim in PROHIBITED_EXACT_CLAIMS:
        for sentence in sentences:
            if claim.lower() in sentence.lower() and not NEGATION_GUARDS.search(sentence):
                hits.append(claim)
                break
    return hits


def required_content_checks(text: str) -> dict[str, bool]:
    normalized = re.sub(r"\s+", " ", text)
    return {
        concept: all(re.search(pattern, normalized, flags=re.IGNORECASE) for pattern in patterns)
        for concept, patterns in REQUIRED_CONTENT_PATTERNS.items()
    }


def render_table_value(value: object) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float):
        return f"{value:.3f}"
    return str(value)


def expected_table_s1(source_path: Path) -> list[list[str]]:
    frame = pd.read_csv(source_path)[
        [
            "cohort",
            "K",
            "seeds_completed",
            "mean_deviance",
            "mean_high_absolute_lag1_fraction",
            "selections_across_three_seeds",
        ]
    ].copy()
    frame["cohort"] = frame["cohort"].map(
        {"mimic": "MIMIC-IV", "eicu": "eICU-CRD", "aumc": "AmsterdamUMCdb"}
    )
    frame.columns = [
        "Dataset",
        "Clusters (K)",
        "Seeds",
        "Mean deviance",
        "Mean high |lag-1| fraction",
        "Selections / 3 seeds",
    ]
    return [list(frame.columns)] + [
        [render_table_value(value) for value in row]
        for row in frame.itertuples(index=False, name=None)
    ]


def compare_table_s1_to_source(table, source_path: Path) -> dict[str, object]:
    expected = expected_table_s1(source_path)
    observed = [
        [cell.text.strip() for cell in row.cells]
        for row in table.rows
    ]
    mismatches: list[dict[str, object]] = []
    maximum_rows = max(len(expected), len(observed))
    for row_index in range(maximum_rows):
        expected_row = expected[row_index] if row_index < len(expected) else []
        observed_row = observed[row_index] if row_index < len(observed) else []
        maximum_columns = max(len(expected_row), len(observed_row))
        for column_index in range(maximum_columns):
            expected_value = (
                expected_row[column_index] if column_index < len(expected_row) else None
            )
            observed_value = (
                observed_row[column_index] if column_index < len(observed_row) else None
            )
            if expected_value != observed_value:
                mismatches.append(
                    {
                        "row": row_index,
                        "column": column_index,
                        "expected": expected_value,
                        "observed": observed_value,
                    }
                )
    return {
        "source_exists": source_path.exists(),
        "expected_rows": len(expected),
        "observed_rows": len(observed),
        "expected_cells": sum(len(row) for row in expected),
        "observed_cells": sum(len(row) for row in observed),
        "mismatch_count": len(mismatches),
        "mismatches": mismatches,
        "exact_match": not mismatches and len(expected) == len(observed),
    }


def figure_asset_identity(
    document: Document,
    caption_prefix: str,
    source_path: Path,
) -> dict[str, object]:
    caption_index = next(
        (
            index
            for index, paragraph in enumerate(document.paragraphs)
            if paragraph.text.strip().startswith(caption_prefix)
        ),
        None,
    )
    result: dict[str, object] = {
        "caption_found": caption_index is not None,
        "source_exists": source_path.exists(),
        "embedded_asset_found": False,
        "identity_method": None,
        "exact_match": False,
    }
    if caption_index is None or not source_path.exists():
        return result

    embedded_blob: bytes | None = None
    embedded_relationship: str | None = None
    for paragraph in reversed(document.paragraphs[:caption_index]):
        blips = paragraph._p.xpath(".//a:blip")
        if not blips:
            continue
        embedded_relationship = blips[-1].get(qn("r:embed"))
        if embedded_relationship in document.part.related_parts:
            embedded_blob = document.part.related_parts[embedded_relationship].blob
        break
    if embedded_blob is None:
        return result

    source_blob = source_path.read_bytes()
    source_sha256 = hashlib.sha256(source_blob).hexdigest()
    embedded_sha256 = hashlib.sha256(embedded_blob).hexdigest()
    result.update(
        {
            "embedded_asset_found": True,
            "embedded_relationship": embedded_relationship,
            "source_sha256": source_sha256,
            "embedded_sha256": embedded_sha256,
        }
    )
    if source_sha256 == embedded_sha256:
        result["identity_method"] = "sha256"
        result["exact_match"] = True
        return result

    try:
        source_image = Image.open(io.BytesIO(source_blob)).convert("RGBA")
        embedded_image = Image.open(io.BytesIO(embedded_blob)).convert("RGBA")
        dimensions_match = source_image.size == embedded_image.size
        pixels_match = dimensions_match and ImageChops.difference(
            source_image, embedded_image
        ).getbbox() is None
        result.update(
            {
                "identity_method": "pixel_comparison",
                "source_dimensions": list(source_image.size),
                "embedded_dimensions": list(embedded_image.size),
                "pixels_match": pixels_match,
                "exact_match": pixels_match,
            }
        )
    except Exception as exc:  # pragma: no cover - defensive reporting path
        result["pixel_comparison_error"] = str(exc)
    return result


def supplement_checks(
    path: Path,
    fresh_k_summary: Path | None = None,
    figure_s2_source: Path | None = None,
    figure_sources: dict[str, Path] | None = None,
) -> dict[str, object]:
    document = Document(path)
    figure_s2_captions = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if re.match(r"^Figure S2(?:\.|\s)", paragraph.text.strip())
    ]
    table_s1 = document.tables[0]
    k_values: list[int] = []
    for row in table_s1.rows[1:]:
        value = row.cells[1].text.strip()
        if value:
            k_values.append(int(float(value)))
    dataset_values = [row.cells[0].text.strip() for row in table_s1.rows[1:]]
    headers = [cell.text.strip() for cell in table_s1.rows[0].cells]
    selection_column = next(
        (index for index, header in enumerate(headers) if "Selections" in header),
        None,
    )
    selection_annotations: dict[str, int] = {}
    if selection_column is not None:
        for row in table_s1.rows[1:]:
            cohort = row.cells[0].text.strip()
            k_value = row.cells[1].text.strip()
            selection = row.cells[selection_column].text.strip()
            if cohort and k_value and selection:
                selection_annotations[f"{cohort}:K{int(float(k_value))}"] = int(float(selection))
    expected_k = [2, 3, 4, 5] * 3
    expected_datasets = {"MIMIC-IV", "eICU-CRD", "AmsterdamUMCdb"}
    expected_selection_keys = {
        f"{cohort}:K{k}" for cohort in expected_datasets for k in (2, 3, 4, 5)
    }
    selection_annotation_present = (
        set(selection_annotations) == expected_selection_keys
        and all(0 <= value <= 3 for value in selection_annotations.values())
    )
    k3_counts = [
        selection_annotations.get(f"{cohort}:K3") for cohort in expected_datasets
    ]
    result: dict[str, object] = {
        "figure_s2_caption_count": len(figure_s2_captions),
        "figure_s2_captions": figure_s2_captions,
        "table_s1_k_values": k_values,
        "table_s1_dataset_values": dataset_values,
        "table_s1_has_only_k2_k5": sorted(k_values) == sorted(expected_k),
        "table_s1_has_three_cohorts": set(dataset_values) == expected_datasets,
        "table_s1_rows": len(table_s1.rows) - 1,
        "table_s1_selection_annotations": selection_annotations,
        "table_s1_selection_annotation_present": selection_annotation_present,
        "table_s1_eicu_k3_not_unanimous": selection_annotations.get("eICU-CRD:K3") != 3,
        "table_s1_k3_not_unanimous_across_cohorts": not all(value == 3 for value in k3_counts),
    }
    if fresh_k_summary is not None:
        if fresh_k_summary.exists():
            result["table_s1_source_identity"] = compare_table_s1_to_source(
                table_s1, fresh_k_summary
            )
        else:
            result["table_s1_source_identity"] = {
                "source_exists": False,
                "exact_match": False,
                "mismatch_count": None,
            }
    resolved_sources = dict(figure_sources or {})
    if figure_s2_source is not None:
        resolved_sources["figure_s2"] = figure_s2_source
    if resolved_sources:
        identities = {
            figure_id: figure_asset_identity(
                document,
                SUPPLEMENT_FIGURE_SOURCES[figure_id][0],
                source_path,
            )
            for figure_id, source_path in resolved_sources.items()
        }
        result["figure_asset_identities"] = identities
        if "figure_s2" in identities:
            result["figure_s2_asset_identity"] = identities["figure_s2"]
    return result


def main() -> int:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    missing = [name for name in REQUIRED_FILES if not (args.package_dir / name).exists()]
    checks: dict[str, object] = {"missing_required_files": missing}
    failures: list[str] = []
    if missing:
        failures.append("Missing required revision documents")

    supplement_path = args.package_dir / "JTIM_Revised_Supplementary_material.docx"
    if supplement_path.exists():
        supplement = supplement_checks(
            supplement_path,
            fresh_k_summary=args.fresh_k_summary,
            figure_sources={
                **default_figure_sources(args.report_root),
                "figure_s2": args.figure_s2_source,
            },
        )
        checks["supplement"] = supplement
        if supplement["figure_s2_caption_count"] != 1:
            failures.append("Figure S2 must have exactly one caption")
        if not supplement["table_s1_has_only_k2_k5"]:
            failures.append("Table S1 must contain exactly K=2-5 for all three cohorts")
        if not supplement["table_s1_has_three_cohorts"]:
            failures.append("Table S1 cohort set is incomplete")
        if not supplement["table_s1_selection_annotation_present"]:
            failures.append("Table S1 must include a complete selections-per-three-seeds annotation")
        if not supplement["table_s1_eicu_k3_not_unanimous"]:
            failures.append("Table S1 must not mark eICU K=3 as unanimously selected")
        if not supplement["table_s1_k3_not_unanimous_across_cohorts"]:
            failures.append("Table S1 must not imply unanimous K=3 selection across cohorts")
        if not supplement.get("table_s1_source_identity", {}).get("exact_match", False):
            failures.append("Table S1 must match fresh_k_grid_summary.csv cell for cell")
        figure_identities = supplement.get("figure_asset_identities", {})
        for figure_id in SUPPLEMENT_FIGURE_SOURCES:
            if not figure_identities.get(figure_id, {}).get("exact_match", False):
                label = SUPPLEMENT_FIGURE_SOURCES[figure_id][0].rstrip(".")
                failures.append(
                    f"Embedded {label} must match the released source figure asset"
                )

    wording_hits: dict[str, list[str]] = {}
    for name in WORDING_FILES:
        path = args.package_dir / name
        if not path.exists():
            continue
        text = author_response_text(path) if name == "JTIM_Point-by-point_Response.docx" else document_text(path)
        hits = prohibited_claim_hits(text)
        wording_hits[name] = hits
        if hits:
            failures.append(f"{name}: prohibited legacy claim remains")
    checks["prohibited_wording_hits"] = wording_hits

    required_content: dict[str, dict[str, bool]] = {}
    for name in CONTENT_FILES:
        path = args.package_dir / name
        if not path.exists():
            continue
        text = author_response_text(path) if name == "JTIM_Point-by-point_Response.docx" else document_text(path)
        concept_checks = required_content_checks(text)
        required_content[name] = concept_checks
        missing_concepts = [concept for concept, passed in concept_checks.items() if not passed]
        if missing_concepts:
            failures.append(f"{name}: missing required scientific content: {', '.join(missing_concepts)}")
    checks["required_scientific_content"] = required_content

    status = {
        "overall_status": "PASS_FINAL_REVISION_DOCUMENT_QA" if not failures else "FAIL_FINAL_REVISION_DOCUMENT_QA",
        "failures": failures,
        "checks": checks,
    }
    (args.output_dir / "revision_document_qa_status.json").write_text(
        json.dumps(status, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    embedded_identities = checks.get("supplement", {}).get("figure_asset_identities", {})
    embedded_pass_count = sum(
        identity.get("exact_match", False)
        for identity in embedded_identities.values()
    )
    report = f"""# Final revision document QA

**Status: {status['overall_status']}**

- Required document files present: {not missing}
- Figure S2 caption count: {checks.get('supplement', {}).get('figure_s2_caption_count', 'not checked')}
- Table S1 restricted to traceable K=2–5 across all three cohorts: {checks.get('supplement', {}).get('table_s1_has_only_k2_k5', False)}
- Table S1 includes seed-level selection annotations and does not mark eICU K=3 as unanimous: {checks.get('supplement', {}).get('table_s1_selection_annotation_present', False) and checks.get('supplement', {}).get('table_s1_eicu_k3_not_unanimous', False)}
- Table S1 matches the source CSV cell for cell: {checks.get('supplement', {}).get('table_s1_source_identity', {}).get('exact_match', False)}
- Embedded revised figures matching released assets: {embedded_pass_count}/{len(SUPPLEMENT_FIGURE_SOURCES)}
- Prohibited legacy claims detected: {sum(len(value) for value in wording_hits.values())}
- Required scientific content checks passed: {all(all(values.values()) for values in required_content.values()) if required_content else False}

This fail-closed check covers document structure, author-authored wording, and required
scientific interpretation in the clean/highlighted manuscript and point-by-point
response. It complements, but does not replace, rendered-page visual inspection.
"""
    (args.output_dir / "W8_REVISION_DOCUMENT_QA.md").write_text(report, encoding="utf-8")
    print(json.dumps(status, indent=2, ensure_ascii=False))
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
