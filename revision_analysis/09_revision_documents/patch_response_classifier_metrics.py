#!/usr/bin/env python3
"""Append the source-linked external classifier comparison to the JTIM response."""

from __future__ import annotations

import argparse
import shutil
import tempfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ANCHOR = (
    "We also compared the exact-version replay of the archived AutoGluon "
    "XGBoost_BAG_L2 model"
)
SENTENCE = (
    " On external AmsterdamUMCdb, the archived model versus the simple comparator "
    "yielded macro OvO AUC 0.769 versus 0.769, balanced accuracy 0.528 versus "
    "0.578, macro F1 0.504 versus 0.583, and multiclass Brier score 0.576 versus "
    "0.407, respectively."
)
R1M4_HEADING = "R1.m4 Nonstandard clustering metric"


def save_safely(document: Document, path: Path) -> None:
    """Save a patched DOCX beside the original and replace it atomically."""
    with tempfile.NamedTemporaryFile(
        prefix=f".{path.stem}.", suffix=path.suffix, dir=path.parent, delete=False
    ) as handle:
        temporary = Path(handle.name)
    try:
        document.save(temporary)
        shutil.copystat(path, temporary)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def append_metrics(path: Path) -> str:
    document = Document(path)
    matches = [paragraph for paragraph in document.paragraphs if ANCHOR in paragraph.text]
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one classifier anchor, found {len(matches)}")
    paragraph = matches[0]
    if SENTENCE.strip() in paragraph.text:
        return "already present"

    run = paragraph.add_run(SENTENCE)
    previous_runs = paragraph.runs[:-1]
    if previous_runs and previous_runs[-1]._element.rPr is not None:
        run._element.insert(0, deepcopy(previous_runs[-1]._element.rPr))

    save_safely(document, path)
    return "appended"


def stabilize_r1m4_pagination(path: Path) -> str:
    """Prevent LibreOffice from placing the R1.m4 block above the top margin."""
    document = Document(path)
    paragraphs = document.paragraphs
    matches = [index for index, paragraph in enumerate(paragraphs) if paragraph.text == R1M4_HEADING]
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one R1.m4 heading, found {len(matches)}")
    index = matches[0]
    heading = paragraphs[index]
    reviewer_comment = paragraphs[index + 1]
    already_stable = (
        heading.paragraph_format.page_break_before is True
        and heading.paragraph_format.keep_with_next is True
        and reviewer_comment.paragraph_format.keep_with_next is True
    )
    if already_stable:
        return "already stable"
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.keep_with_next = True
    reviewer_comment.paragraph_format.keep_with_next = True
    save_safely(document, path)
    return "stabilized"


def normalize_page_footer(path: Path) -> str:
    """Rebuild the PAGE field with separate runs so two-digit even pages render."""
    document = Document(path)
    footer = document.sections[0].footer
    paragraph = footer.paragraphs[0]
    field_types = [
        node.get(qn("w:fldCharType"))
        for node in paragraph._p.iter(qn("w:fldChar"))
    ]
    if len(paragraph.runs) >= 5 and field_types == ["begin", "separate", "end"]:
        return "already normalized"

    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def styled_run():
        run = paragraph.add_run()
        run.font.name = "Times New Roman"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:ascii"), "Times New Roman"
        )
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        return run

    begin_run = styled_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._element.append(begin)

    instruction_run = styled_run()
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    instruction_run._element.append(instruction)

    separate_run = styled_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._element.append(separate)

    result_run = styled_run()
    result_run.text = "1"

    end_run = styled_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._element.append(end)

    save_safely(document, path)
    return "normalized"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", type=Path)
    args = parser.parse_args()
    metrics_result = append_metrics(args.path)
    pagination_result = stabilize_r1m4_pagination(args.path)
    footer_result = normalize_page_footer(args.path)
    print(f"Classifier comparison metrics: {metrics_result}")
    print(f"R1.m4 pagination: {pagination_result}")
    print(f"Page footer: {footer_result}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
