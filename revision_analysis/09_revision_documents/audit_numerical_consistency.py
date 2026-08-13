#!/usr/bin/env python3
"""Audit source-to-document numerical consistency for the final JTIM revision."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml.ns import qn


DOCUMENTS = {
    "manuscript_clean": "JTIM_Revised_Manuscript_clean.docx",
    "manuscript_highlight": "JTIM_Revised_Manuscript_highlight.docx",
    "supplement": "JTIM_Revised_Supplementary_material.docx",
    "response": "JTIM_Point-by-point_Response.docx",
}


@dataclass
class Check:
    check_id: str
    category: str
    source: str
    document: str
    location: str
    check_type: str
    expected: str
    observed: str
    passed: bool
    details: str = ""


def parse_args() -> argparse.Namespace:
    analysis = Path(__file__).resolve().parents[2]
    project = analysis.parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--package-dir",
        type=Path,
        default=(
            project
            / "01_manuscript/14、Journal of Translational Internal Medicine/revision_package_20260805"
        ),
    )
    parser.add_argument(
        "--report-root",
        type=Path,
        default=analysis / "02_revision_outputs/reports",
    )
    parser.add_argument(
        "--table-root",
        type=Path,
        default=analysis / "02_revision_outputs/tables/harmonized_longitudinal_tables",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=analysis / "02_revision_outputs/reports/W9_numerical_consistency",
    )
    return parser.parse_args()


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def paragraph_with_anchor(document: Document, anchor: str) -> str:
    anchor_lower = anchor.lower()
    for paragraph in document.paragraphs:
        text = normalize(paragraph.text)
        if anchor_lower in text.lower():
            return text
    return ""


def check_anchor_tokens(
    checks: list[Check],
    *,
    check_id: str,
    category: str,
    source: str,
    document_name: str,
    document: Document,
    anchor: str,
    tokens: list[str],
    location: str,
) -> None:
    observed = paragraph_with_anchor(document, anchor)
    missing = [token for token in tokens if token not in observed]
    checks.append(
        Check(
            check_id=check_id,
            category=category,
            source=source,
            document=document_name,
            location=location,
            check_type="anchored_text_tokens",
            expected=json.dumps(tokens, ensure_ascii=False),
            observed=observed,
            passed=bool(observed) and not missing,
            details=(
                "all expected source-derived tokens present"
                if observed and not missing
                else f"missing tokens: {missing}"
            ),
        )
    )


def render_value(value: object) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float):
        return f"{value:.3f}"
    return str(value)


def frame_grid(frame: pd.DataFrame) -> list[list[str]]:
    return [list(frame.columns)] + [
        [render_value(value) for value in row]
        for row in frame.itertuples(index=False, name=None)
    ]


def table_grid(table) -> list[list[str]]:
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def find_tables(document: Document, headers: list[str]) -> list:
    return [
        table
        for table in document.tables
        if [cell.text.strip() for cell in table.rows[0].cells] == headers
    ]


def grid_mismatches(
    expected: list[list[str]], observed: list[list[str]]
) -> list[dict[str, object]]:
    mismatches: list[dict[str, object]] = []
    for row_index in range(max(len(expected), len(observed))):
        expected_row = expected[row_index] if row_index < len(expected) else []
        observed_row = observed[row_index] if row_index < len(observed) else []
        for column_index in range(max(len(expected_row), len(observed_row))):
            expected_value = (
                expected_row[column_index] if column_index < len(expected_row) else None
            )
            observed_value = (
                observed_row[column_index] if column_index < len(observed_row) else None
            )
            if expected_value != observed_value:
                mismatches.append(
                    {
                        "row": row_index,
                        "column": column_index,
                        "expected": expected_value,
                        "observed": observed_value,
                    }
                )
    return mismatches


def check_dataframe_table(
    checks: list[Check],
    *,
    check_id: str,
    category: str,
    source: str,
    document_name: str,
    document: Document,
    location: str,
    frame: pd.DataFrame,
    occurrence: int = 0,
) -> None:
    expected = frame_grid(frame)
    matching_tables = find_tables(document, list(frame.columns))
    table = matching_tables[occurrence] if occurrence < len(matching_tables) else None
    observed = table_grid(table) if table is not None else []
    mismatches = grid_mismatches(expected, observed)
    checks.append(
        Check(
            check_id=check_id,
            category=category,
            source=source,
            document=document_name,
            location=location,
            check_type="table_cell_identity",
            expected=f"{len(expected)} rows; {sum(len(row) for row in expected)} cells",
            observed=f"{len(observed)} rows; {sum(len(row) for row in observed)} cells",
            passed=table is not None and not mismatches,
            details=json.dumps(
                {
                    "matching_tables": len(matching_tables),
                    "selected_occurrence": occurrence,
                    "mismatch_count": len(mismatches),
                    "mismatches": mismatches[:20],
                },
                ensure_ascii=False,
            ),
        )
    )


def embedded_figure_blob(document: Document, caption_prefix: str) -> bytes | None:
    caption_index = next(
        (
            index
            for index, paragraph in enumerate(document.paragraphs)
            if paragraph.text.strip().startswith(caption_prefix)
        ),
        None,
    )
    if caption_index is None:
        return None
    for paragraph in reversed(document.paragraphs[:caption_index]):
        blips = paragraph._p.xpath(".//a:blip")
        if not blips:
            continue
        relationship = blips[-1].get(qn("r:embed"))
        if relationship in document.part.related_parts:
            return document.part.related_parts[relationship].blob
        return None
    return None


def check_figure_identity(
    checks: list[Check],
    *,
    check_id: str,
    category: str,
    source_path: Path,
    source_label: str,
    document_name: str,
    document: Document,
    caption_prefix: str,
) -> None:
    embedded = embedded_figure_blob(document, caption_prefix)
    source = source_path.read_bytes() if source_path.exists() else None
    source_hash = hashlib.sha256(source).hexdigest() if source is not None else ""
    embedded_hash = hashlib.sha256(embedded).hexdigest() if embedded is not None else ""
    checks.append(
        Check(
            check_id=check_id,
            category=category,
            source=source_label,
            document=document_name,
            location=caption_prefix.rstrip("."),
            check_type="embedded_asset_sha256",
            expected=source_hash,
            observed=embedded_hash,
            passed=bool(source_hash) and source_hash == embedded_hash,
            details="exact byte identity between released source figure and DOCX media",
        )
    )


def table_s1_frame(report_root: Path) -> pd.DataFrame:
    frame = pd.read_csv(report_root / "W3_fresh_k_grid/fresh_k_grid_summary.csv")[[
        "cohort",
        "K",
        "seeds_completed",
        "mean_deviance",
        "mean_high_absolute_lag1_fraction",
        "selections_across_three_seeds",
    ]].copy()
    frame["cohort"] = frame["cohort"].map(
        {"mimic": "MIMIC-IV", "eicu": "eICU-CRD", "aumc": "AmsterdamUMCdb"}
    )
    frame.columns = [
        "Dataset",
        "Clusters (K)",
        "Seeds",
        "Mean deviance",
        "Mean high |lag-1| fraction",
        "Selections / 3 seeds",
    ]
    return frame


def supplementary_frames(report_root: Path) -> list[tuple[str, str, str, pd.DataFrame]]:
    fresh = pd.read_csv(report_root / "W3_fresh_k_grid/fresh_k_grid_summary.csv")[[
        "cohort", "K", "seeds_completed", "mean_deviance",
        "mean_high_absolute_lag1_fraction", "mean_uncertain_fraction",
        "median_max_posterior_probability", "median_selection_score",
        "selections_across_three_seeds",
    ]].rename(columns={
        "cohort": "Cohort", "seeds_completed": "Seeds",
        "mean_deviance": "Mean deviance",
        "mean_high_absolute_lag1_fraction": "Mean high |lag-1| fraction",
        "mean_uncertain_fraction": "Mean uncertain fraction",
        "median_max_posterior_probability": "Median maximum posterior",
        "median_selection_score": "Median screening score",
        "selections_across_three_seeds": "Selections / 3 seeds",
    })
    deep = pd.read_csv(report_root / "W3_deep_k_stability/deep_k_cohort_summary.csv").rename(columns={
        "cohort": "Cohort", "k3_fits": "K=3 starts",
        "k3_fits_passing": "K=3 starts passing",
        "median_agreement_vs_archived": "Median agreement vs archived",
        "minimum_ari_vs_archived": "Minimum ARI vs archived",
        "minimum_cluster_prevalence": "Minimum cluster prevalence",
        "maximum_high_lag1_fraction": "Maximum high |lag-1| fraction",
        "minimum_pairwise_agreement": "Minimum cross-start agreement",
        "minimum_pairwise_ari": "Minimum cross-start ARI",
        "k3_cross_seed_status": "K=3 stability status",
        "K2_selections": "K=2 selections / 3", "K3_selections": "K=3 selections / 3",
    })
    deep_robustness = pd.read_csv(
        report_root / "W3_deep_robustness/deep_robustness_scenario_summary.csv"
    ).rename(columns={
        "cohort": "Cohort", "scenario": "Sensitivity scenario",
        "seeds_completed": "Starts", "passing_seeds": "Passing starts",
        "median_exact_agreement": "Median agreement",
        "minimum_exact_agreement": "Minimum agreement", "median_ari": "Median ARI",
        "minimum_ari": "Minimum ARI", "median_nmi": "Median NMI",
        "maximum_uncertain_fraction": "Maximum uncertain fraction",
        "minimum_cluster_prevalence": "Minimum cluster prevalence",
        "median_dr_fraction": "Median DR fraction",
        "median_rr_fraction": "Median RR fraction",
        "median_pw_fraction": "Median PW fraction",
        "maximum_high_lag1_fraction": "Maximum high |lag-1| fraction",
        "scenario_status": "Deep-rerun status",
    })
    robustness = pd.read_csv(
        report_root / "W3_cross_cohort_robustness/cross_cohort_robustness_metrics.csv"
    )[[
        "cohort", "scenario", "refit_patients", "retained_fraction",
        "exact_agreement", "ari", "nmi", "uncertain_fraction",
        "minimum_cluster_prevalence", "screening_threshold_status",
    ]].rename(columns={
        "cohort": "Cohort", "scenario": "Sensitivity scenario",
        "refit_patients": "Patients", "retained_fraction": "Retained fraction",
        "exact_agreement": "Exact agreement", "ari": "ARI", "nmi": "NMI",
        "uncertain_fraction": "Uncertain fraction",
        "minimum_cluster_prevalence": "Minimum cluster prevalence",
        "screening_threshold_status": "Threshold status",
    })
    urine = pd.read_csv(
        report_root / "W3_uo_multiseed_sensitivity/uo_multiseed_scenario_summary.csv"
    ).rename(columns={
        "scenario": "Scenario", "seeds_completed": "Starts",
        "passing_seeds": "Passing starts", "patients": "Patients",
        "retained_fraction": "Retained fraction",
        "median_exact_agreement": "Median agreement",
        "minimum_exact_agreement": "Minimum agreement", "median_ari": "Median ARI",
        "minimum_ari": "Minimum ARI", "median_nmi": "Median NMI",
        "maximum_uncertain_fraction": "Maximum uncertain fraction",
        "minimum_cluster_prevalence": "Minimum cluster prevalence",
        "maximum_high_lag1_fraction": "Maximum high |lag-1| fraction",
        "scenario_status": "Status",
    })
    archived = pd.read_csv(
        report_root / "W4_classifier_validation/archived_xgboost_bag_l2_metrics.csv"
    )[[
        "cohort", "model", "OVO_macro_AUC_test", "OVO_macro_AUC_val",
        "accuracy", "precision_macro", "recall_macro", "f1_macro",
    ]].rename(columns={
        "cohort": "Cohort", "model": "Model",
        "OVO_macro_AUC_test": "Macro OvO AUC (test)",
        "OVO_macro_AUC_val": "Macro OvO AUC (validation)",
        "accuracy": "Accuracy", "precision_macro": "Macro precision",
        "recall_macro": "Macro recall", "f1_macro": "Macro F1",
    })
    comparator = pd.read_csv(
        report_root / "W4_classifier_validation/primary_and_comparator_metrics.csv"
    )[[
        "model", "cohort", "n", "balanced_accuracy", "f1_macro",
        "auc_ovo_macro", "multiclass_brier",
    ]].rename(columns={
        "model": "Model", "cohort": "Cohort", "n": "N",
        "balanced_accuracy": "Balanced accuracy", "f1_macro": "Macro F1",
        "auc_ovo_macro": "Macro OvO AUC", "multiclass_brier": "Multiclass Brier",
    })
    per_class = pd.read_csv(
        report_root / "W4_classifier_validation/revalidation_class_metrics.csv"
    ).rename(columns={
        "cohort": "Cohort", "class": "Class", "phenotype": "Phenotype",
        "support": "N", "prevalence": "Prevalence", "precision": "Precision",
        "sensitivity_recall": "Sensitivity", "specificity": "Specificity",
        "npv": "NPV", "f1": "F1", "auc_ovr": "OvR AUC",
    })
    mortality = pd.read_csv(
        report_root / "W5_independent_outcomes/outcome_adjusted_effects.csv"
    )
    mortality = mortality[
        mortality.outcome.eq("mortality_28d")
        & mortality.analysis_population.eq("overall")
        & mortality.model_variant.eq("primary_onset_nonrenal_sofa")
    ][[
        "cohort_label", "comparison", "adjusted_or", "ci_low", "ci_high",
        "p_value", "n_complete", "events",
    ]].copy()
    mortality["p_value"] = mortality["p_value"].map(
        lambda value: "<0.001" if value < 0.001 else f"{value:.3f}"
    )
    mortality = mortality.rename(columns={
        "cohort_label": "Cohort", "comparison": "Comparison",
        "adjusted_or": "Adjusted OR", "ci_low": "95% CI lower",
        "ci_high": "95% CI upper", "p_value": "P value",
        "n_complete": "N", "events": "Events",
    })
    diuretic = pd.read_csv(
        report_root / "W6_diuretic_exploratory/early_first_dose_pooled_models.csv"
    )
    diuretic = diuretic[
        diuretic.response_definition.isin(
            ["response_archived_first", "response_strict_10pct_200"]
        )
    ].copy()
    diuretic["response_definition"] = diuretic["response_definition"].map({
        "response_archived_first": "Archived first-day threshold",
        "response_strict_10pct_200": "Strict ≥10% and ≥200 mL",
    })
    diuretic["p_value"] = diuretic["p_value"].map(
        lambda value: "<0.001" if value < 0.001 else f"{value:.3f}"
    )
    diuretic = diuretic[[
        "cohort_label", "response_definition",
        "adjusted_or_response_vs_nonresponse", "ci_low", "ci_high", "p_value",
        "n_complete", "deaths", "responsive_n",
    ]].rename(columns={
        "cohort_label": "Cohort", "response_definition": "Definition",
        "adjusted_or_response_vs_nonresponse": "Adjusted OR",
        "ci_low": "95% CI lower", "ci_high": "95% CI upper",
        "p_value": "P value", "n_complete": "N", "deaths": "Deaths",
        "responsive_n": "Responders",
    })
    return [
        ("S15d", "cluster_selection", "W3_fresh_k_grid/fresh_k_grid_summary.csv", fresh),
        ("S15e", "cluster_selection", "W3_deep_k_stability/deep_k_cohort_summary.csv", deep),
        ("S15f", "robustness", "W3_deep_robustness/deep_robustness_scenario_summary.csv", deep_robustness),
        ("S15g", "robustness", "W3_cross_cohort_robustness/cross_cohort_robustness_metrics.csv", robustness),
        ("S16", "urine_output_sensitivity", "W3_uo_multiseed_sensitivity/uo_multiseed_scenario_summary.csv", urine),
        ("S17a", "classifier", "W4_classifier_validation/archived_xgboost_bag_l2_metrics.csv", archived),
        ("S17b", "classifier", "W4_classifier_validation/primary_and_comparator_metrics.csv", comparator),
        ("S17c", "classifier", "W4_classifier_validation/revalidation_class_metrics.csv", per_class),
        ("S18", "mortality", "W5_independent_outcomes/outcome_adjusted_effects.csv", mortality),
        ("S19", "diuretic", "W6_diuretic_exploratory/early_first_dose_pooled_models.csv", diuretic),
    ]


def check_main_table1(
    checks: list[Check], documents: dict[str, Document], report_root: Path
) -> None:
    counts = pd.read_csv(report_root / "W2_cross_cohort_inputs/cluster_label_counts.csv")
    expected = ["Patient number"]
    for cohort in ["MIMIC-IV", "AUMC", "eICU-CRD"]:
        subset = counts[counts.cohort.eq(cohort)].set_index("phenotype")
        expected.extend(
            [
                str(int(subset.patients.sum())),
                str(int(subset.loc["RR", "patients"])),
                str(int(subset.loc["DR", "patients"])),
                str(int(subset.loc["PW", "patients"])),
                "",
            ]
        )
    for key in ["manuscript_clean", "manuscript_highlight"]:
        document = documents[key]
        observed = []
        if document.tables:
            for row in document.tables[0].rows:
                values = [cell.text.strip() for cell in row.cells]
                if values and values[0] == "Patient number":
                    observed = values
                    break
        mismatches = grid_mismatches([expected], [observed] if observed else [])
        checks.append(
            Check(
                check_id=f"{key}.table1_patient_counts",
                category="population",
                source="W2_cross_cohort_inputs/cluster_label_counts.csv",
                document=DOCUMENTS[key],
                location="Table 1 — Patient number",
                check_type="table_row_identity",
                expected=json.dumps(expected, ensure_ascii=False),
                observed=json.dumps(observed, ensure_ascii=False),
                passed=bool(observed) and not mismatches,
                details=f"mismatch_count={len(mismatches)}",
            )
        )


def longitudinal_frames(table_root: Path) -> list[tuple[str, pd.DataFrame]]:
    mapping = [
        ("S3", table_root / "Table_S3_MIMIC-IV.csv"),
        ("S4", table_root / "Table_S4_AUMC.csv"),
        ("S5", table_root / "Table_S5_eICU-CRD.csv"),
    ]
    frames: list[tuple[str, pd.DataFrame]] = []
    for table_id, path in mapping:
        raw = pd.read_csv(path, header=None)
        days = raw.iloc[1, 2:].astype(str).tolist()
        values = raw.iloc[3:].copy()
        values.columns = ["Feature", "Unit"] + [f"d{x}" for x in days]
        for phenotype, start in [("RR", 2), ("DR", 10), ("PW", 18)]:
            panel = pd.concat([values.iloc[:, :2], values.iloc[:, start : start + 8]], axis=1)
            panel.columns = [
                "Feature", "Unit", "Day −1", "Day 1", "Day 2", "Day 3",
                "Day 4", "Day 5", "Day 6", "Day 7",
            ]
            frames.append((f"{table_id}_{phenotype}_part1", panel.iloc[:20]))
            frames.append((f"{table_id}_{phenotype}_part2", panel.iloc[20:]))
    return frames


def add_text_checks(
    checks: list[Check], documents: dict[str, Document], report_root: Path
) -> None:
    inventory = pd.read_csv(report_root / "W2_cross_cohort_inputs/cluster_matrix_inventory.csv")
    patient_map = dict(zip(inventory.cohort, inventory.patients))
    total = int(inventory.patients.sum())
    population_tokens = [
        f"{total:,} patients",
        f"{int(patient_map['MIMIC-IV']):,}",
        f"{int(patient_map['AUMC']):,}",
        f"{int(patient_map['eICU-CRD']):,}",
    ]
    labels = pd.read_csv(report_root / "W2_cross_cohort_inputs/cluster_label_counts.csv")
    phenotype_tokens = []
    for cohort, label in [
        ("MIMIC-IV", "MIMIC-IV"),
        ("AUMC", "AmsterdamUMCdb"),
        ("eICU-CRD", "eICU-CRD"),
    ]:
        subset = labels[labels.cohort.eq(cohort)].set_index("phenotype")
        phenotype_tokens.append(
            f"{label} {subset.loc['RR', 'fraction']:.1%}, "
            f"{subset.loc['DR', 'fraction']:.1%}, and {subset.loc['PW', 'fraction']:.1%}"
        )

    fresh = pd.read_csv(report_root / "W3_fresh_k_grid/fresh_k_grid_summary.csv")
    selection = fresh[fresh.K.eq(3)].set_index("cohort").selections_across_three_seeds
    k_token = (
        f"k=3 was selected in {int(selection['mimic'])}/3 MIMIC-IV, "
        f"{int(selection['eicu'])}/3 eICU-CRD, and "
        f"{int(selection['aumc'])}/3 AmsterdamUMCdb"
    )
    uo = pd.read_csv(
        report_root / "W3_uo_multiseed_sensitivity/uo_multiseed_scenario_summary.csv"
    ).set_index("scenario")
    documented = uo.loc["documented_windows"]
    coverage = uo.loc["high_coverage"]
    uo_tokens = [
        f"{int(documented.patients):,} patients",
        f"median agreement {documented.median_exact_agreement:.3f}",
        f"median ARI {documented.median_ari:.3f}",
        f"{int(coverage.patients):,} patients",
        f"{coverage.maximum_high_lag1_fraction:.1%}",
    ]

    mortality = pd.read_csv(report_root / "W5_independent_outcomes/outcome_adjusted_effects.csv")
    mortality = mortality[
        mortality.outcome.eq("mortality_28d")
        & mortality.analysis_population.eq("overall")
        & mortality.model_variant.eq("primary_onset_nonrenal_sofa")
    ]
    mortality_tokens = [
        f"{row.adjusted_or:.2f} (95% CI {row.ci_low:.2f}–{row.ci_high:.2f})"
        for row in mortality.itertuples()
    ]
    compact_tokens = []
    for cohort in ["MIMIC-IV", "eICU-CRD", "AUMC"]:
        subset = mortality[mortality.cohort_label.eq(cohort)].set_index("comparison")
        compact_tokens.append(
            f"{cohort} DR/PW {subset.loc['DR vs RR', 'adjusted_or']:.2f}/"
            f"{subset.loc['PW vs RR', 'adjusted_or']:.2f}"
        )

    diuretic = pd.read_csv(
        report_root / "W6_diuretic_exploratory/early_first_dose_pooled_models.csv"
    )
    archived_diuretic = diuretic[
        diuretic.response_definition.eq("response_archived_first")
    ].set_index("cohort_label")
    interactions = pd.read_csv(
        report_root / "W6_diuretic_exploratory/early_first_dose_interaction_tests.csv"
    ).set_index("cohort_label")
    diuretic_main_tokens = []
    diuretic_response_tokens = []
    for cohort in ["MIMIC-IV", "AUMC"]:
        row = archived_diuretic.loc[cohort]
        diuretic_response_tokens.extend(
            [f"{row.adjusted_or_response_vs_nonresponse:.3f}", f"{row.ci_low:.3f}–{row.ci_high:.3f}"]
        )
        diuretic_main_tokens.extend(
            [
                f"{row.adjusted_or_response_vs_nonresponse:.3f}",
                f"{row.ci_low:.3f}–{row.ci_high:.3f}",
                f"p={row.p_value:.3f}",
            ]
        )
    diuretic_main_tokens.extend(
        [f"p={interactions.loc['MIMIC-IV', 'p_value']:.3f}", f"p={interactions.loc['AUMC', 'p_value']:.3f}"]
    )

    archived = pd.read_csv(
        report_root / "W4_classifier_validation/archived_xgboost_bag_l2_metrics.csv"
    ).set_index("cohort")
    pairwise = pd.read_csv(
        report_root / "W4_classifier_validation/primary_and_comparator_pairwise_auc.csv"
    )
    pairwise = pairwise[pairwise.model.eq("archived_AutoGluon_XGBoost_BAG_L2")]
    classifier_tokens = []
    for cohort in ["internal_MIMIC_eICU", "external_AUMC"]:
        row = archived.loc[cohort]
        classifier_tokens.extend(
            [f"{row.OVO_macro_AUC_test:.3f}", f"{row.accuracy:.3f}", f"{row.f1_macro:.3f}"]
        )
        classifier_tokens.extend(
            f"{value:.3f}"
            for value in pairwise[pairwise.cohort.eq(cohort)].auc_pair_normalized
        )
    bootstrap = pd.read_csv(
        report_root / "W4_classifier_validation/revalidation_bootstrap_ci.csv"
    )
    for cohort in ["internal_MIMIC_eICU", "external_AUMC"]:
        row = bootstrap[
            bootstrap.cohort.eq(cohort) & bootstrap.metric.eq("auc_ovo_macro")
        ].iloc[0]
        classifier_tokens.extend([f"{row.ci_low:.3f}", f"{row.ci_high:.3f}"])

    comparator = pd.read_csv(
        report_root / "W4_classifier_validation/primary_and_comparator_metrics.csv"
    )
    external = comparator[comparator.cohort.eq("external_AUMC")].set_index("model")
    response_classifier_tokens = []
    for model in ["archived_AutoGluon_XGBoost_BAG_L2", "simple_six_variable_logistic"]:
        row = external.loc[model]
        response_classifier_tokens.extend(
            [
                f"{row.auc_ovo_macro:.3f}",
                f"{row.balanced_accuracy:.3f}",
                f"{row.f1_macro:.3f}",
                f"{row.multiclass_brier:.3f}",
            ]
        )

    for key in ["manuscript_clean", "manuscript_highlight"]:
        document = documents[key]
        check_anchor_tokens(
            checks, check_id=f"{key}.population", category="population",
            source="W2_cross_cohort_inputs/cluster_matrix_inventory.csv",
            document_name=DOCUMENTS[key], document=document,
            anchor="The study included", tokens=population_tokens,
            location="Results — Study population",
        )
        check_anchor_tokens(
            checks, check_id=f"{key}.phenotype_distribution", category="population",
            source="W2_cross_cohort_inputs/cluster_label_counts.csv",
            document_name=DOCUMENTS[key], document=document,
            anchor="In the reproducible archived eICU candidate set",
            tokens=phenotype_tokens, location="Results — Cluster selection and distribution",
        )
        check_anchor_tokens(
            checks, check_id=f"{key}.k_and_uo", category="cluster_robustness",
            source="W3_fresh_k_grid; W3_deep_k_stability; W3_uo_multiseed_sensitivity",
            document_name=DOCUMENTS[key], document=document,
            anchor="In the reproducible archived eICU candidate set",
            tokens=[k_token, *uo_tokens], location="Results — Cluster robustness",
        )
        check_anchor_tokens(
            checks, check_id=f"{key}.mortality", category="mortality",
            source="W5_independent_outcomes/outcome_adjusted_effects.csv",
            document_name=DOCUMENTS[key], document=document,
            anchor="Twenty-eight-day mortality remained", tokens=mortality_tokens,
            location="Results — Independent mortality association",
        )
        check_anchor_tokens(
            checks, check_id=f"{key}.diuretic", category="diuretic",
            source="W6_diuretic_exploratory/early_first_dose_pooled_models.csv; early_first_dose_interaction_tests.csv",
            document_name=DOCUMENTS[key], document=document,
            anchor="In the restricted first-dose", tokens=diuretic_main_tokens,
            location="Results — Exploratory diuretic response",
        )
        check_anchor_tokens(
            checks, check_id=f"{key}.classifier", category="classifier",
            source="W4_classifier_validation aggregate CSVs",
            document_name=DOCUMENTS[key], document=document,
            anchor="The archived XGBoost model", tokens=classifier_tokens,
            location="Results — Classifier",
        )

    response = documents["response"]
    check_anchor_tokens(
        checks, check_id="response.k_uo_primary", category="cluster_robustness",
        source="W3_fresh_k_grid; W3_deep_k_stability",
        document_name=DOCUMENTS["response"], document=response,
        anchor="Cluster number and initialization.", tokens=[k_token],
        location="E.2 response",
    )
    check_anchor_tokens(
        checks, check_id="response.uo_repeated", category="urine_output_sensitivity",
        source="W3_uo_multiseed_sensitivity/uo_multiseed_scenario_summary.csv",
        document_name=DOCUMENTS["response"], document=response,
        anchor="The documented-window sensitivity retained", tokens=uo_tokens,
        location="R1.4 response",
    )
    check_anchor_tokens(
        checks, check_id="response.mortality", category="mortality",
        source="W5_independent_outcomes/outcome_adjusted_effects.csv",
        document_name=DOCUMENTS["response"], document=response,
        anchor="Relative to RR, adjusted ORs for DR and PW were", tokens=mortality_tokens,
        location="E.3 response",
    )
    check_anchor_tokens(
        checks, check_id="response.diuretic", category="diuretic",
        source="W6_diuretic_exploratory/early_first_dose_pooled_models.csv",
        document_name=DOCUMENTS["response"], document=response,
        anchor="In a restricted first-dose 24-hour landmark", tokens=diuretic_response_tokens,
        location="E.4 response",
    )
    check_anchor_tokens(
        checks, check_id="response.classifier", category="classifier",
        source="W4_classifier_validation/primary_and_comparator_metrics.csv",
        document_name=DOCUMENTS["response"], document=response,
        anchor="We also compared the exact-version replay",
        tokens=response_classifier_tokens, location="E.5 response",
    )

    incremental = pd.read_csv(
        report_root / "W4_classifier_validation/paired_incremental_value.csv"
    )
    external_incremental = incremental[incremental.cohort.eq("external_AUMC")].set_index("metric")
    interpretation_pass = (
        external_incremental.loc["auc_ovo_macro", "ci_low"] <= 0
        <= external_incremental.loc["auc_ovo_macro", "ci_high"]
        and external_incremental.loc["balanced_accuracy", "ci_high"] < 0
        and external_incremental.loc["f1_macro", "ci_high"] < 0
        and external_incremental.loc["multiclass_brier", "ci_low"] > 0
        and "no external incremental value" in paragraph_with_anchor(
            documents["manuscript_clean"], "The archived XGBoost model"
        )
    )
    checks.append(
        Check(
            check_id="classifier.incremental_value_interpretation",
            category="classifier",
            source="W4_classifier_validation/paired_incremental_value.csv",
            document=DOCUMENTS["manuscript_clean"],
            location="Results — Classifier",
            check_type="source_ci_interpretation",
            expected="external AUC CI crosses 0; balanced accuracy/F1 worse; Brier worse",
            observed=external_incremental[["difference_autogluon_minus_simple", "ci_low", "ci_high"]].to_json(),
            passed=bool(interpretation_pass),
            details="validates the narrative claim of no external incremental advantage",
        )
    )


def main() -> int:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    missing_documents = [
        filename for filename in DOCUMENTS.values() if not (args.package_dir / filename).exists()
    ]
    if missing_documents:
        raise FileNotFoundError(f"Missing final revision documents: {missing_documents}")
    documents = {
        key: Document(args.package_dir / filename) for key, filename in DOCUMENTS.items()
    }
    checks: list[Check] = []

    add_text_checks(checks, documents, args.report_root)
    check_main_table1(checks, documents, args.report_root)

    supplement = documents["supplement"]
    check_dataframe_table(
        checks, check_id="supplement.table_s1", category="cluster_selection",
        source="W3_fresh_k_grid/fresh_k_grid_summary.csv",
        document_name=DOCUMENTS["supplement"], document=supplement,
        location="Table S1", frame=table_s1_frame(args.report_root),
    )
    for table_id, category, source, frame in supplementary_frames(args.report_root):
        check_dataframe_table(
            checks, check_id=f"supplement.table_{table_id.lower()}", category=category,
            source=source, document_name=DOCUMENTS["supplement"], document=supplement,
            location=f"Table {table_id}", frame=frame,
        )

    longitudinal_sources = {
        "S3": "Table_S3_MIMIC-IV.csv",
        "S4": "Table_S4_AUMC.csv",
        "S5": "Table_S5_eICU-CRD.csv",
    }
    for occurrence, (table_id, frame) in enumerate(longitudinal_frames(args.table_root)):
        source_table = table_id.split("_")[0]
        check_dataframe_table(
            checks, check_id=f"supplement.table_{table_id.lower()}", category="longitudinal_tables",
            source=f"harmonized_longitudinal_tables/{longitudinal_sources[source_table]}",
            document_name=DOCUMENTS["supplement"], document=supplement,
            location=f"Corrected {table_id}", frame=frame, occurrence=occurrence,
        )

    figure_sources = [
        (
            "figure_s2", "cluster_selection",
            args.report_root / "W3_deep_k_stability/Figure_S2_cross_cohort_k_stability.png",
            "W3_deep_k_stability/Figure_S2_cross_cohort_k_stability.png", "Figure S2.",
        ),
        (
            "figure_s10", "robustness",
            args.report_root / "W3_cross_cohort_robustness/cross_cohort_robustness_matrix.png",
            "W3_cross_cohort_robustness/cross_cohort_robustness_matrix.png", "Figure S10.",
        ),
        (
            "figure_s11a", "missingness_sensitivity",
            args.report_root / "W3_cluster_robustness/documented_windows_cluster_sensitivity.png",
            "W3_cluster_robustness/documented_windows_cluster_sensitivity.png", "Figure S11a.",
        ),
        (
            "figure_s11b", "missingness_sensitivity",
            args.report_root / "W3_cluster_robustness/high_coverage_cluster_sensitivity.png",
            "W3_cluster_robustness/high_coverage_cluster_sensitivity.png", "Figure S11b.",
        ),
        (
            "figure_s12", "classifier",
            args.report_root / "W4_classifier_validation/Figure_S12_archived_model_calibration_comparator.png",
            "W4_classifier_validation/Figure_S12_archived_model_calibration_comparator.png", "Figure S12.",
        ),
        (
            "figure_s13", "landmark_outcomes",
            args.report_root / "W5_independent_outcomes/W5_adjusted_outcomes_forest.png",
            "W5_independent_outcomes/W5_adjusted_outcomes_forest.png", "Figure S13.",
        ),
        (
            "figure_s14", "diuretic",
            args.report_root / "W6_diuretic_exploratory/W6_early_diuretic_response_forest.png",
            "W6_diuretic_exploratory/W6_early_diuretic_response_forest.png", "Figure S14.",
        ),
    ]
    for check_id, category, source_path, source_label, caption in figure_sources:
        check_figure_identity(
            checks, check_id=f"supplement.{check_id}", category=category,
            source_path=source_path, source_label=source_label,
            document_name=DOCUMENTS["supplement"], document=supplement,
            caption_prefix=caption,
        )

    clean_numbers = re.findall(r"(?<![A-Za-z])\d[\d,]*(?:\.\d+)?%?", normalize(
        " ".join(paragraph.text for paragraph in documents["manuscript_clean"].paragraphs)
        + " "
        + " ".join(cell.text for table in documents["manuscript_clean"].tables for row in table.rows for cell in row.cells)
    ))
    highlight_numbers = re.findall(r"(?<![A-Za-z])\d[\d,]*(?:\.\d+)?%?", normalize(
        " ".join(paragraph.text for paragraph in documents["manuscript_highlight"].paragraphs)
        + " "
        + " ".join(cell.text for table in documents["manuscript_highlight"].tables for row in table.rows for cell in row.cells)
    ))
    checks.append(
        Check(
            check_id="manuscript.clean_vs_highlight_numeric_sequence",
            category="cross_document_identity",
            source="JTIM_Revised_Manuscript_clean.docx",
            document=DOCUMENTS["manuscript_highlight"],
            location="whole document",
            check_type="numeric_token_sequence_identity",
            expected=f"{len(clean_numbers)} numeric tokens",
            observed=f"{len(highlight_numbers)} numeric tokens",
            passed=clean_numbers == highlight_numbers,
            details="clean and highlighted manuscripts must contain the same numeric sequence",
        )
    )

    failures = [check for check in checks if not check.passed]
    status = {
        "overall_status": (
            "PASS_FINAL_NUMERICAL_CONSISTENCY_AUDIT"
            if not failures
            else "FAIL_FINAL_NUMERICAL_CONSISTENCY_AUDIT"
        ),
        "checks_total": len(checks),
        "checks_passed": len(checks) - len(failures),
        "checks_failed": len(failures),
        "failed_check_ids": [check.check_id for check in failures],
        "category_summary": {
            category: {
                "total": sum(check.category == category for check in checks),
                "passed": sum(check.category == category and check.passed for check in checks),
            }
            for category in sorted({check.category for check in checks})
        },
        "scope_note": (
            "The audit verifies source-linked revision claims and regenerated tables. "
            "Legacy phenotype-characterizing recovery tables without a revision aggregate "
            "source are outside this automated source-to-location gate."
        ),
    }
    (args.output_dir / "numerical_consistency_status.json").write_text(
        json.dumps(status, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    pd.DataFrame([asdict(check) for check in checks]).to_csv(
        args.output_dir / "numerical_consistency_checks.csv", index=False
    )
    report_lines = [
        "# Final numerical consistency audit",
        "",
        f"**Status: {status['overall_status']}**",
        "",
        f"- Checks passed: {status['checks_passed']} / {status['checks_total']}",
        f"- Failed check IDs: {', '.join(status['failed_check_ids']) or 'None'}",
        "- Documents: clean manuscript, highlighted manuscript, supplement, and point-by-point response",
        "- Sources: aggregate revision reports and regenerated longitudinal tables",
        "",
        "## Category summary",
        "",
        "| Category | Passed | Total |",
        "|---|---:|---:|",
    ]
    for category, values in status["category_summary"].items():
        report_lines.append(f"| {category} | {values['passed']} | {values['total']} |")
    report_lines.extend(
        [
            "",
            "## Scope boundary",
            "",
            status["scope_note"],
            "",
            "The audit is fail-closed for mapped source-to-document claims, table cells, "
            "released figure assets, and repeated numeric claims. It does not substitute "
            "for author confirmation of metadata, signatures, or final page/line numbers.",
            "",
        ]
    )
    (args.output_dir / "W9_NUMERICAL_CONSISTENCY_AUDIT.md").write_text(
        "\n".join(report_lines), encoding="utf-8"
    )
    print(json.dumps(status, indent=2, ensure_ascii=False))
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
