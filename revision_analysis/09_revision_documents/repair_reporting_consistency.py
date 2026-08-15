#!/usr/bin/env python3
"""Repair and audit the final Table S2 and Table S19 reporting inconsistencies.

The script is deliberately narrow. It reads the frozen, current eICU aggregate
characteristics export, rejects incompatible denominators or percentages, writes
an identifier-free publication table, and patches only the affected supplement
cells and explanatory text. The broader eICU ``pulmonary`` category is never
substituted for the manuscript's ``chronic pulmonary disease`` row.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from copy import deepcopy
from pathlib import Path

import pandas as pd
from docx import Document


ANALYSIS = Path(__file__).resolve().parents[2]
PROJECT = ANALYSIS.parents[1]
MANUSCRIPT = PROJECT / "01_manuscript/14、Journal of Translational Internal Medicine"
SUPPLEMENT = MANUSCRIPT / "revision_package/JTIM_Revised_Supplementary_material.docx"
TABLE_S2_SOURCE = (
    ANALYSIS
    / "00_frozen_inputs/data_snapshot/remote_project_snapshot/"
    "03.eICU_SAKI_trajCluster/result/eICU_clusters_characteristics.csv"
)
DIURETIC_SOURCE = (
    ANALYSIS
    / "02_revision_outputs/reports/W6_diuretic_exploratory/"
    "early_first_dose_pooled_models.csv"
)
OUTPUT_DIR = ANALYSIS / "02_revision_outputs/reports/W1_table_s2_audit"

GROUP_COLUMNS = {"Overall": "Overall", "RR": "C2", "DR": "C1", "PW": "C3"}
EXPECTED_DENOMINATORS = {"Overall": 1417, "RR": 869, "DR": 423, "PW": 125}
TABLE_S2_ROWS = (
    ("Patient number, N", None, "Current analysed cohort denominators"),
    ("Myocardial infarct, n (%)", "Myocardial infarct, n (%)", ""),
    ("Congestive heart failure, n (%)", "Congestive heart failure, n (%)", ""),
    (
        "Peripheral vascular disease, n (%)",
        None,
        "Compatible eICU diagnosis field unavailable",
    ),
    (
        "Cerebrovascular disease, n (%)",
        None,
        "Compatible eICU diagnosis field unavailable",
    ),
    (
        "Chronic pulmonary disease, n (%)",
        None,
        "Broad eICU pulmonary category is not definitionally compatible and was not substituted",
    ),
    ("Hepatic disease, n (%)", "Hepatic, n (%)", ""),
    ("Diabetes, n (%)", "Diabetes, n (%)", ""),
    ("Hypertension, n (%)", "Hypertension, n (%)", ""),
    ("Respiratory failure, n (%)", "Respiratory failure, n (%)", ""),
    ("Septic shock, n (%)", "Septic shock, n (%)", ""),
    ("Acidosis, n (%)", "Acidosis, n (%)", ""),
)
RESPONSE_LABELS = {
    "response_archived_first": "Archived first-day threshold",
    "response_strict_10pct_200": "Strict ≥10% and ≥200 mL",
    "response_absolute_200": "Absolute ≥200 mL",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_count_percent(value: str) -> tuple[int, float]:
    match = re.fullmatch(r"(\d+) \((\d+(?:\.\d+)?)\)", value.strip())
    if not match:
        raise ValueError(f"Invalid count/percentage cell: {value!r}")
    return int(match.group(1)), float(match.group(2))


def authoritative_table_s2(source: Path) -> pd.DataFrame:
    raw = pd.read_csv(source, dtype=str, keep_default_na=False)
    required = {"feature", "type", "Overall", "C1", "C2", "C3", "P-Value"}
    if set(raw.columns) != required:
        raise ValueError(f"Unexpected eICU characteristics schema: {list(raw.columns)}")

    n_rows = raw.loc[raw.feature.eq("n")]
    if len(n_rows) != 1:
        raise ValueError("Expected exactly one eICU denominator row")
    n_row = n_rows.iloc[0]
    observed_denominators = {
        label: int(n_row[source_column])
        for label, source_column in GROUP_COLUMNS.items()
    }
    if observed_denominators != EXPECTED_DENOMINATORS:
        raise ValueError(
            "Table S2 source is not the authoritative N=1,417 cohort: "
            f"{observed_denominators}"
        )
    if sum(observed_denominators[label] for label in ("RR", "DR", "PW")) != 1417:
        raise ValueError("eICU phenotype denominators do not sum to the overall cohort")

    publication_rows: list[dict[str, str]] = []
    for document_label, source_feature, note in TABLE_S2_ROWS:
        if source_feature is None:
            if document_label == "Patient number, N":
                values = {label: str(value) for label, value in observed_denominators.items()}
                p_value = ""
                source_field = "n"
            else:
                values = {label: "—" for label in GROUP_COLUMNS}
                p_value = "—"
                source_field = "Not available"
        else:
            feature_rows = raw.loc[raw.feature.eq(source_feature)]
            selected = feature_rows.loc[feature_rows.type.eq("1")]
            if len(selected) != 1:
                raise ValueError(f"Missing unique positive-indicator row: {source_feature}")
            source_row = selected.iloc[0]
            values = {
                label: source_row[source_column]
                for label, source_column in GROUP_COLUMNS.items()
            }
            for label, value in values.items():
                count, percentage = parse_count_percent(value)
                denominator = observed_denominators[label]
                if count > denominator:
                    raise ValueError(
                        f"Count exceeds denominator for {source_feature}/{label}: "
                        f"{count}>{denominator}"
                    )
                expected_percentage = 100.0 * count / denominator
                if abs(percentage - expected_percentage) > 0.051:
                    raise ValueError(
                        f"Percentage/denominator mismatch for {source_feature}/{label}: "
                        f"reported {percentage:.1f}, expected {expected_percentage:.1f}"
                    )
            p_values = [value for value in feature_rows["P-Value"] if value]
            if len(p_values) != 1:
                raise ValueError(f"Missing unique P value for {source_feature}")
            p_value = p_values[0]
            source_field = source_feature

        publication_rows.append(
            {
                "Characteristic": document_label,
                **values,
                "P value": p_value,
                "Source field": source_field,
                "Availability note": note,
            }
        )

    result = pd.DataFrame(publication_rows)
    if "pulmonary, n (%)" in set(result["Source field"]):
        raise ValueError("Broad pulmonary category must not populate chronic pulmonary disease")
    return result


def diuretic_table_s19(source: Path) -> pd.DataFrame:
    raw = pd.read_csv(source)
    selected = raw.loc[
        raw.response_definition.isin(RESPONSE_LABELS),
        [
            "cohort_label",
            "response_definition",
            "adjusted_or_response_vs_nonresponse",
            "ci_low",
            "ci_high",
            "p_value",
            "n_complete",
            "deaths",
            "responsive_n",
        ],
    ].copy()
    expected = {
        (cohort, definition)
        for cohort in ("MIMIC-IV", "AUMC")
        for definition in RESPONSE_LABELS
    }
    observed = set(zip(selected.cohort_label, selected.response_definition))
    if observed != expected:
        raise ValueError(f"Incomplete Table S19 response-definition grid: {expected - observed}")
    selected["cohort_order"] = selected.cohort_label.map({"MIMIC-IV": 0, "AUMC": 1})
    selected["definition_order"] = selected.response_definition.map(
        {definition: index for index, definition in enumerate(RESPONSE_LABELS)}
    )
    selected = selected.sort_values(["cohort_order", "definition_order"])

    rows = []
    for row in selected.itertuples():
        rows.append(
            {
                "Cohort": row.cohort_label,
                "Definition": RESPONSE_LABELS[row.response_definition],
                "Adjusted OR": f"{row.adjusted_or_response_vs_nonresponse:.3f}",
                "95% CI lower": f"{row.ci_low:.3f}",
                "95% CI upper": f"{row.ci_high:.3f}",
                "P value": "<0.001" if row.p_value < 0.001 else f"{row.p_value:.3f}",
                "N": str(int(row.n_complete)),
                "Deaths": str(int(row.deaths)),
                "Responders": str(int(row.responsive_n)),
            }
        )
    return pd.DataFrame(rows)


def replace_text_preserving_format(container, text: str) -> None:
    """Replace paragraph or cell text while retaining its first-run formatting."""
    paragraphs = container.paragraphs if hasattr(container, "paragraphs") else [container]
    first = paragraphs[0]
    if first.runs:
        first.runs[0].text = text
        for run in first.runs[1:]:
            run.text = ""
    else:
        first.add_run(text)
    for paragraph in paragraphs[1:]:
        for run in paragraph.runs:
            run.text = ""


def find_table_s2(document: Document):
    matches = []
    for table in document.tables:
        if len(table.rows) < 3 or len(table.columns) != 12:
            continue
        if table.rows[2].cells[0].text.strip() == "Patient number, N":
            matches.append(table)
    if len(matches) != 1:
        raise ValueError(f"Expected one Table S2, found {len(matches)}")
    return matches[0]


def find_table_by_header(document: Document, header: list[str]):
    matches = [
        table
        for table in document.tables
        if [cell.text.strip() for cell in table.rows[0].cells] == header
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one table with header {header}, found {len(matches)}")
    return matches[0]


def patch_table_s2(document: Document, frame: pd.DataFrame) -> None:
    table = find_table_s2(document)
    row_map = {row.cells[0].text.strip(): row for row in table.rows[2:]}
    for source_row in frame.itertuples(index=False):
        label = source_row[0]
        if label not in row_map:
            raise ValueError(f"Table S2 row missing from document: {label}")
        values = list(source_row[1:6])
        for cell, value in zip(row_map[label].cells[7:12], values):
            replace_text_preserving_format(cell, str(value))

    note = (
        "Data are presented as No. (%). The AmsterdamUMCdb (AUMC) dataset is excluded "
        "from this table because compatible ICD diagnosis codes were unavailable. "
        "For eICU-CRD, em dashes indicate that compatible peripheral vascular, "
        "cerebrovascular, or chronic pulmonary disease variables were unavailable; the "
        "broader eICU pulmonary diagnosis category was not substituted for chronic "
        "pulmonary disease. P-values were calculated using the Chi-square test or "
        "Fisher's exact test. Abbreviations: RR, Rapid Recovery; DR, Delayed Recovery; "
        "PW, Progressive Worsening."
    )
    paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith("Data are presented as No. (%)")
    ]
    if len(paragraphs) != 1:
        raise ValueError(f"Expected one Table S2 note paragraph, found {len(paragraphs)}")
    replace_text_preserving_format(paragraphs[0], note)


def patch_integrity_audit(document: Document) -> None:
    paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith("Automated source-to-output comparison")
    ]
    if len(paragraphs) != 1:
        raise ValueError(f"Expected one data-integrity paragraph, found {len(paragraphs)}")
    text = (
        "Automated source-to-output comparison found zero mismatches between the verified "
        "source matrix and generated cohort Tables S3–S5. The embedded MIMIC-IV workbook "
        "contained 22 mismatches confined to RR day 7 (glucose through pH), following the "
        "preceding correct FiO2 value as a sequential +1 spreadsheet autofill; that "
        "workbook was not an input to clustering, outcome, diuretic, or classifier "
        "analyses. Final quality control also found that the eICU-CRD block of Table S2 "
        "retained comorbidity counts from a legacy approximately 1,970-patient export "
        "beside the current N=1,417 denominators. We regenerated that descriptive block "
        "from the authoritative N=1,417 aggregate. The incompatible broad pulmonary "
        "category was not relabelled as chronic pulmonary disease. Table S2 is descriptive "
        "and was not an input to clustering, outcome, diuretic, or classifier analyses. "
        "Unit audit also identified table-layer harmonization requirements for eICU-CRD "
        "FiO2, AmsterdamUMCdb bilirubin, and AmsterdamUMCdb hematocrit; Tables S3–S5 were "
        "regenerated after cohort-specific conversion and range checks."
    )
    replace_text_preserving_format(paragraphs[0], text)

    table = find_table_by_header(
        document,
        [
            "Cohort",
            "Source→generated mismatches",
            "Generated→embedded mismatches",
            "Action",
        ],
    )
    existing = next(
        (row for row in table.rows[1:] if row.cells[0].text.strip() == "eICU-CRD Table S2"),
        None,
    )
    if existing is None:
        new_row = deepcopy(table.rows[-1]._tr)
        table._tbl.append(new_row)
        existing = table.rows[-1]
    values = [
        "eICU-CRD Table S2",
        "0 after regeneration",
        "Not applicable",
        "Legacy denominator block replaced; incompatible pulmonary field not substituted",
    ]
    for cell, value in zip(existing.cells, values):
        replace_text_preserving_format(cell, value)


def patch_table_s19(document: Document, frame: pd.DataFrame) -> None:
    header = list(frame.columns)
    table = find_table_by_header(document, header)
    templates = [deepcopy(row._tr) for row in table.rows[1:3]]
    if len(templates) < 2:
        raise ValueError("Table S19 lacks styled data-row templates")
    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)
    for row_index, values in enumerate(frame.itertuples(index=False, name=None)):
        table._tbl.append(deepcopy(templates[row_index % 2]))
        output_row = table.rows[-1]
        for cell, value in zip(output_row.cells, values):
            replace_text_preserving_format(cell, str(value))

    captions = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith("Figure S14.")
    ]
    if len(captions) != 1:
        raise ValueError(f"Expected one Figure S14 caption, found {len(captions)}")
    caption = (
        "Figure S14. Restricted first-dose landmark associations between diuretic "
        "response and 28-day mortality. Each cohort shows the archived first-day "
        "threshold, the strict ≥10% and ≥200 mL definition, and the absolute ≥200 mL "
        "definition. Points are adjusted odds ratios for response versus nonresponse and "
        "bars are 95% confidence intervals. These treated-only, post-treatment "
        "associations are non-causal and hypothesis-generating."
    )
    replace_text_preserving_format(captions[0], caption)


def write_table_s2_evidence(frame: pd.DataFrame, source: Path, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    frame.to_csv(output_dir / "eicu_table_s2_verified.csv", index=False)
    status = {
        "overall_status": "PASS_TABLE_S2_AUTHORITY_AUDIT",
        "authoritative_source_sha256": sha256(source),
        "authoritative_overall_n": 1417,
        "phenotype_denominators": {"RR": 869, "DR": 423, "PW": 125},
        "broad_pulmonary_substituted_for_chronic_pulmonary": False,
        "publication_rows": len(frame),
    }
    (output_dir / "table_s2_audit_status.json").write_text(
        json.dumps(status, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    report = "\n".join(
        [
            "# eICU Table S2 authority audit",
            "",
            "**Status: PASS_TABLE_S2_AUTHORITY_AUDIT**",
            "",
            "- Current analysed cohort: N=1,417 (RR 869; DR 423; PW 125).",
            "- Every reported count is within its current denominator and every percentage was recomputed against that denominator.",
            "- Peripheral vascular, cerebrovascular, and chronic pulmonary disease remain unavailable in eICU-CRD.",
            "- The broader eICU pulmonary diagnosis category was not substituted for chronic pulmonary disease.",
            "- Table S2 is descriptive and is not an input to clustering, outcomes, diuretic, or classifier analyses.",
            "",
        ]
    )
    (output_dir / "W1_TABLE_S2_AUTHORITY_AUDIT.md").write_text(report, encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--supplement", type=Path, default=SUPPLEMENT)
    parser.add_argument("--table-s2-source", type=Path, default=TABLE_S2_SOURCE)
    parser.add_argument("--diuretic-source", type=Path, default=DIURETIC_SOURCE)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    parser.add_argument("--check-only", action="store_true")
    args = parser.parse_args()

    table_s2 = authoritative_table_s2(args.table_s2_source)
    table_s19 = diuretic_table_s19(args.diuretic_source)
    write_table_s2_evidence(table_s2, args.table_s2_source, args.output_dir)
    if not args.check_only:
        document = Document(args.supplement)
        patch_table_s2(document, table_s2)
        patch_integrity_audit(document)
        patch_table_s19(document, table_s19)
        document.save(args.supplement)
    print(
        json.dumps(
            {
                "supplement": str(args.supplement),
                "table_s2_rows": len(table_s2),
                "table_s19_rows": len(table_s19),
                "check_only": args.check_only,
            },
            indent=2,
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
