"""Aggregate-only window figure, supplement text and author review document."""
import json
from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

REPO=Path(__file__).resolve().parents[3]
OUT=REPO/'02_revision_outputs/reports/20260921_window_prediction'
m=pd.read_csv(OUT/'metrics.csv');ci=pd.read_csv(OUT/'auc_confidence_intervals.csv');delta=pd.read_csv(OUT/'paired_changes_vs24h.csv')
manifest=json.loads((OUT/'input_manifest.json').read_text())
assert len(m)==12 and len(ci)==48 and len(delta)==40
def row(h,s):return m[(m.hours==h)&(m.split==s)].iloc[0]
def val(h,s,k):return float(row(h,s)[k])
def ci_text(h,s,k):
    r=ci[(ci.hours==h)&(ci.split==s)&(ci.metric==k)].iloc[0]
    return f'{r.auc:.3f} ({r.lower95:.3f}–{r.upper95:.3f})'

methods=('We compared cumulative observation windows of 6, 12, 18, 24, 30, and 36 hours after SA-AKI onset while retaining the original subphenotype labels and training, internal-test, and external-test partitions (4,903, 1,227, and 2,183 patients, respectively). '
 'For each window, maximum, minimum, mean, and pairwise differences were derived from the same 27 variables. '
 'Predictors were reconstructed from archived forward-filled trajectories before cohort-wide multiple imputation; only observations available by each cutoff were used, and remaining missing values were handled by XGBoost. '
 'The XGBoost hyperparameters selected in the preceding 24-hour training-set comparison were held fixed across windows. Separate models were fitted using identical five-fold partitions, with no parameter or window selection based on the external test set. '
 'Models were fitted with AutoGluon 1.6.3 using a maximum tree depth of 3, learning rate of 0.03, row and column sampling fractions of 0.85, minimum child weight of 1, L2 regularization of 1, and up to 1,000 boosting rounds with fold-specific early stopping. '
 'We reported macro one-versus-one AUC and all three pairwise AUCs, using pair-normalized probabilities for pairwise comparisons. Pointwise 95% confidence intervals and paired differences relative to 24 hours were obtained from 1,000 class-stratified bootstrap resamples of each test set, conditional on the fitted models.')
results=(f'At 24 hours, macro OvO AUC was {val(24,"internal","auc_ovo_macro"):.3f} internally and {val(24,"external","auc_ovo_macro"):.3f} externally; DR–PW AUC was {val(24,"internal","auc_DR_PW"):.3f} and {val(24,"external","auc_DR_PW"):.3f}, respectively. '
 f'At 36 hours, the corresponding macro OvO AUCs were {val(36,"internal","auc_ovo_macro"):.3f} and {val(36,"external","auc_ovo_macro"):.3f}, and DR–PW AUCs were {val(36,"internal","auc_DR_PW"):.3f} and {val(36,"external","auc_DR_PW"):.3f}.')
change=delta[(delta.hours==36)&(delta.split=='external')&(delta.metric=='auc_DR_PW')].iloc[0]
results+=f' The external DR–PW AUC difference between 36 and 24 hours was {change.delta_auc:+.3f} (95% CI, {change.lower95:+.3f} to {change.upper95:+.3f}).'
caption=('Prediction performance across observation windows. Panels show DR versus PW, overall macro OvO AUC, DR versus RR, and RR versus PW. '
 'Lines show point estimates in the internal test set (n = 1,227) and external AUMC test set (n = 2,183); shaded areas show pointwise 95% confidence intervals from 1,000 class-stratified bootstrap resamples. '
 'The dashed vertical line marks the 24-hour reference window. Original labels, test populations, model hyperparameters, and training fold partitions were held fixed. '
 'DR, Delayed Recovery; RR, Rapid Recovery; PW, Progressive Worsening; OvO, one-versus-one.')
response=(f'To examine how the available observation period affected discrimination, we evaluated six cumulative windows from 6 to 36 hours using the original subphenotype labels and identical train–test partitions. We retained the same XGBoost hyperparameters across windows and reported overall and all three pairwise AUCs in both test cohorts. '
 f'External DR–PW AUC was {val(24,"external","auc_DR_PW"):.3f} at 24 hours and {val(36,"external","auc_DR_PW"):.3f} at 36 hours (paired difference, {change.delta_auc:+.3f}; 95% CI, {change.lower95:+.3f} to {change.upper95:+.3f}). '
 'The 24-hour analysis remains the primary early-classification analysis; the later windows quantify the change in discrimination with a longer observation period.')

# Figure contract: quantitative grid; compare all four discrimination metrics,
# with DR-PW first. Fixed populations, full six windows, paired-bootstrap source.
plt.rcParams.update({'font.family':'sans-serif','font.sans-serif':['Arial','DejaVu Sans'],
 'font.size':8,'axes.titlesize':9,'axes.labelsize':8,'legend.fontsize':8,
 'axes.spines.top':False,'axes.spines.right':False,'pdf.fonttype':42,'svg.fonttype':'none'})
fig,axs=plt.subplots(2,2,figsize=(7.2,5.8),sharex=True,sharey=True)
for ax,key,title,letter in zip(axs.flat,['auc_DR_PW','auc_ovo_macro','auc_DR_RR','auc_RR_PW'],
                             ['DR versus PW','Overall macro OvO','DR versus RR','RR versus PW'],'abcd'):
    for split,color,marker,label in [('internal','#37668C','o','Internal'),('external','#B6693E','s','External AUMC')]:
        z=ci[(ci.metric==key)&(ci.split==split)].sort_values('hours')
        ax.fill_between(z.hours.to_numpy(),z.lower95.to_numpy(),z.upper95.to_numpy(),color=color,alpha=.12,lw=0)
        ax.plot(z.hours,z.auc,color=color,marker=marker,ms=3.8,lw=1.5,label=label)
    ax.axvline(24,color='#888888',ls='--',lw=.8,zorder=0)
    ax.set_title(title,loc='left',pad=7);ax.text(-.15,1.06,letter,transform=ax.transAxes,fontweight='bold',fontsize=11)
    ax.set_xticks([6,12,18,24,30,36]);ax.set_ylim(.5,1.0);ax.set_yticks(np.arange(.5,1.01,.1))
    ax.grid(axis='y',color='#eeeeee',lw=.5)
for ax in axs[:,0]:ax.set_ylabel('AUC')
for ax in axs[-1,:]:ax.set_xlabel('Hours after SA-AKI onset')
handles,labels=axs[0,0].get_legend_handles_labels();fig.legend(handles,labels,loc='upper center',ncol=2,bbox_to_anchor=(.54,1.01),frameon=False)
fig.subplots_adjust(left=.1,right=.98,bottom=.1,top=.91,wspace=.22,hspace=.36)
for ext in ['pdf','svg','png','tiff']:
    fig.savefig(OUT/f'window_prediction.{ext}',dpi=300 if ext=='png' else 600,bbox_inches='tight',**({'pil_kwargs':{'compression':'tiff_lzw'}} if ext=='tiff' else {}))
plt.close(fig)

table=[]
for s in ['internal','external']:
 for h in [6,12,18,24,30,36]:
    table.append({'Cohort':s,'Window (h)':h,'Overall macro OvO':ci_text(h,s,'auc_ovo_macro'),
                  'DR vs RR':ci_text(h,s,'auc_DR_RR'),'DR vs PW':ci_text(h,s,'auc_DR_PW'),'RR vs PW':ci_text(h,s,'auc_RR_PW')})
pd.DataFrame(table).to_csv(OUT/'supplementary_table_auc_95CI.csv',index=False)

overview=['# 六个观察窗口的预测结果','',
 '本轮使用原论文分组和原患者划分，完成了 6、12、18、24、30、36 小时的固定参数 XGBoost 比较。模型、预测结果及运行材料已备份到本地私有目录；活动正文和补充材料未改动。','',
 '| 窗口 | 内部总体 AUC | 外部总体 AUC | 内部 DR–PW | 外部 DR–PW |','|---|---:|---:|---:|---:|']
for h in [6,12,18,24,30,36]:overview.append(f'| {h} h | {val(h,"internal","auc_ovo_macro"):.3f} | {val(h,"external","auc_ovo_macro"):.3f} | {val(h,"internal","auc_DR_PW"):.3f} | {val(h,"external","auc_DR_PW"):.3f} |')
overview+=['',f'外部 DR–PW 从 24 到 36 小时的 AUC 变化为 {change.delta_auc:+.3f}，配对 bootstrap 95% CI 为 {change.lower95:+.3f} 至 {change.upper95:+.3f}。',
 '', '## 本轮与上一轮 24 小时结果的关系','',
 '旧轨迹表经过完整随访队列的多重插补后才截取观察窗口。本轮从此前保存的前向填充表重建，保留剩余缺失值，由 XGBoost 处理，避免窗口外观测参与后续多重插补。原标签、患者、基线肌酐、单位和分组均未改变；基线肌酐重算的原始比值与归档比值完全一致。',
 '因此本轮 24 小时结果与 6–36 小时系列使用同一预处理口径；不要把上一轮外部 AUC 0.781 拼入本轮曲线。前一轮调参结果作为参数来源和历史比较保存。',
 '', '## 如何用于修稿','',
 '建议保留 24 小时作为主时间点；将完整六窗口图和表作为一项补充分析。较长窗口的性能变化不能替代 24 小时性能，也不能直接证明临床决策获益。',
 '若确认采用本轮 24 小时模型作为正文主模型，必须同步分类器方法、摘要与 Results 数字、Figure 5 和对应性能表；旧八算法比较及旧 SHAP 图不能直接代表新模型。SHAP 尚未重算，须在主模型确认后同步。聚类、Table 1 及结局分析无需因此重做。',
 '', '## 完成与核对','',
 '- 六个窗口均为原训练 4,903 人、内部测试 1,227 人、外部 AUMC 2,183 人；窗口间患者、标签及顺序一致。',
 '- 固定 XGBoost 参数，五折、一次重复、无堆叠；实际折数及每折成员哈希一致。',
 '- 18 个队列与窗口组合均通过未来数据扰动不变性检查。该检查验证特征构建的时间截断，不替代所有历史上游数据处理的审计。',
 '- 12 组测试结果由保存的预测概率独立重算；总体和成对 AUC 公式交叉校验一致。',
 '- 1,000 次按类别分层的配对 bootstrap，区间反映固定模型在测试样本上的不确定性。',
 '', '## 文件','',
 '- `window_prediction.pdf` / `.svg` / `.tiff`：图件；`.png`：预览。',
 '- `supplementary_table_auc_95CI.csv`：全部总体及成对 AUC 和区间。',
 '- `metrics.csv`：完整性能指标；`paired_changes_vs24h.csv`：各窗口相对 24 小时的配对变化。',
 '- `补充分析与回复建议.docx`：独立审阅附件；`01_英文材料草案.md`：可复制的英文方法、结果、图注及回复草案。',
 '- `private_backup/`：受限模型及患者级预测备份，不属于投稿或公开交付材料。']
(OUT/'00_结果说明.md').write_text('\n'.join(overview)+'\n')
english='# Proposed supplementary material and response\n\nAuthor review draft. These passages have not been inserted into the active manuscript, supplement, or response. Figure and table numbering will be assigned after approval.\n\n## Supplementary methods\n\n'+methods+'\n\n## Supplementary results\n\n'+results+'\n\n## Figure legend\n\n'+caption+'\n\n## Proposed response addition\n\n'+response+'\n'
(OUT/'01_英文材料草案.md').write_text(english)

# New author-review attachment, not an edit to the tracked manuscript.
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
 rp=st.element.get_or_add_rPr()
 for f in list(rp.findall(qn('w:rFonts'))):rp.remove(f)
 fonts=OxmlElement('w:rFonts')
 for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('cs','Times New Roman'),('eastAsia','Noto Sans CJK SC')]:fonts.set(qn('w:'+k),v)
 rp.insert(0,fonts)
 for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
 st.font.size=Pt(11 if name=='Normal' else 15 if name=='Title' else 12)
 st.paragraph_format.line_spacing=1.25;st.paragraph_format.space_after=Pt(6)
doc.add_paragraph('观察窗口预测分析审阅材料',style='Title')
doc.add_paragraph('包含六个窗口的图表、补充方法与结果，以及回复建议。原稿尚未修改；图表编号及主模型采用方案待确认。')
doc.add_paragraph('本轮结果',style='Heading 1');doc.add_paragraph(results)
doc.add_paragraph('与上一轮的区别',style='Heading 1')
doc.add_paragraph('本轮在完整队列多重插补之前取数，各时点只使用截至该时点的观测，保留剩余缺失值供 XGBoost 处理。六个窗口使用同一流程，24 小时结果也已重跑。原患者、分组标签和训练测试划分不变。')
doc.add_paragraph('Supplementary methods',style='Heading 1');doc.add_paragraph(methods)
doc.add_page_break();doc.add_paragraph('Observation window comparison',style='Heading 1')
doc.add_picture(str(OUT/'window_prediction.png'),width=Cm(17));doc.add_paragraph(caption)
doc.add_page_break();doc.add_paragraph('AUC across observation windows',style='Heading 1')
for s,label in [('internal','Internal test set n = 1227'),('external','External AUMC test set n = 2183')]:
 doc.add_paragraph(label,style='Heading 2')
 t=doc.add_table(rows=1,cols=5);t.autofit=False
 headers=['Hours','Macro OvO','DR vs RR','DR vs PW','RR vs PW'];widths=[1.5,3.875,3.875,3.875,3.875]
 for cell,h,w in zip(t.rows[0].cells,headers,widths):cell.text=h;cell.width=Cm(w)
 for h in [6,12,18,24,30,36]:
  cells=t.add_row().cells
  for c,v,w in zip(cells,[str(h)]+[ci_text(h,s,k).replace(' (','\n(') for k in ['auc_ovo_macro','auc_DR_RR','auc_DR_PW','auc_RR_PW']],widths):c.text=v;c.width=Cm(w)
 for ri,tr in enumerate(t.rows):
  for cell in tr.cells:
   tcPr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single' if (ri==0 and side in ['top','bottom']) or (ri==6 and side=='bottom') else 'nil');e.set(qn('w:sz'),'6');b.append(e)
   tcPr.append(b)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1
    for run in p.runs:run.font.size=Pt(9);run.bold=ri==0
 doc.add_paragraph('')
doc.add_paragraph('Values are AUC (95% CI). All windows use the same test populations; confidence intervals are based on 1,000 class-stratified bootstrap resamples.')
doc.add_page_break();doc.add_paragraph('Proposed response addition',style='Heading 1');doc.add_paragraph(response)
doc.add_paragraph('正文采用前需要同步的内容',style='Heading 1')
doc.add_paragraph('若采用本轮 24 小时模型，需要同步摘要和正文的分类器数字、方法中的训练方案、Figure 5 及性能表，并重算对应 SHAP。现有八算法比较与旧 SHAP 不能直接作为新模型结果。原稿保持不变，本附件中的文字为待确认建议。')
for p in doc.paragraphs:
 for run in p.runs:
  rp=run._r.get_or_add_rPr()
  fonts=rp.find(qn('w:rFonts'))
  if fonts is None:fonts=OxmlElement('w:rFonts');rp.insert(0,fonts)
  for k in list(fonts.attrib):del fonts.attrib[k]
  for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','Noto Sans CJK SC')]:fonts.set(qn('w:'+k),v)
  run.font.color.rgb=RGBColor(0,0,0)
doc.save(OUT/'补充分析与回复建议.docx')
print('Created figure, tables, English draft and author review DOCX')
