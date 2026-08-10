from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from docx import Document


MODULE_PATH = (
    Path(__file__).resolve().parents[1]
    / "09_revision_documents"
    / "audit_reviewer_traceability.py"
)
SPEC = importlib.util.spec_from_file_location("audit_reviewer_traceability", MODULE_PATH)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


def test_response_section_parser_preserves_complete_sections():
    sections = MODULE.parse_sections(
        "# Response\n\n## E.1 First\n\n> comment\n\n**Response:** one\n\n"
        "## E.2 Second\n\n> comment\n\n**Response:** two\n"
    )
    assert set(sections) == {"E.1 First", "E.2 Second"}
    assert "**Response:** one" in sections["E.1 First"]
    assert "**Response:** two" in sections["E.2 Second"]


def test_expected_comment_identifiers_are_unique_and_complete():
    titles = [requirement.title for requirement in MODULE.REQUIREMENTS]
    assert len(titles) == len(set(titles)) == 15
    assert titles[0].startswith("E.1 ")
    assert titles[-1].startswith("R1.m5 ")


def test_token_match_is_case_insensitive_and_fail_closed():
    assert MODULE.contains_all("No External Incremental Value", ("external incremental",))[0]
    passed, missing = MODULE.contains_all("No external value", ("incremental",))
    assert passed is False
    assert missing == ["incremental"]


def test_final_word_sections_are_split_and_preserve_required_content(tmp_path):
    path = tmp_path / "response.docx"
    doc = Document()
    doc.add_heading("E.1 First", level=1)
    doc.add_paragraph("Original reviewer comment")
    doc.add_paragraph("Response: zero mismatches and reporting artifact")
    doc.add_paragraph("Changes in the manuscript: Tables S3–S5 and data dictionary")
    doc.add_heading("E.2 Second", level=1)
    doc.add_paragraph("Second original comment")
    doc.add_paragraph("Response: one of three and material sensitivity")
    doc.add_paragraph("Changes in the manuscript: Figure S2")
    doc.save(path)

    sections = MODULE.parse_docx_sections(path, {"E.1 First", "E.2 Second"})
    assert set(sections) == {"E.1 First", "E.2 Second"}
    assert "zero mismatches" in sections["E.1 First"]
    assert "Figure S2" in sections["E.2 Second"]


def test_upstream_qa_summary_is_dynamic_and_fail_closed():
    w8 = {
        "overall_status": "PASS_FINAL_REVISION_DOCUMENT_QA",
        "checks": {"supplement": {
            "table_s1_source_identity": {"expected_cells": 78, "mismatch_count": 0},
            "figure_asset_identities": {
                "s2": {"exact_match": True},
                "s10": {"exact_match": True},
            },
        }},
    }
    w9 = {
        "overall_status": "PASS_FINAL_NUMERICAL_CONSISTENCY_AUDIT",
        "checks_passed": 57,
        "checks_total": 57,
    }
    w11 = {
        "overall_status": "PASS",
        "artifact_pass_count": 28,
        "artifact_count": 28,
        "visual_review": {
            "s2": {"status": "PASS", "hash_match": True},
            "s10": {"status": "PASS", "hash_match": True},
        },
    }
    summary = MODULE.summarize_upstream_qa(w8, w9, w11)
    assert summary["w9_checks_passed"] == 57
    assert summary["all_pass"] is True

    w9["checks_passed"] = 56
    assert MODULE.summarize_upstream_qa(w8, w9, w11)["all_pass"] is False
