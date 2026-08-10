from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from PIL import Image


MODULE_PATH = (
    Path(__file__).resolve().parents[1]
    / "09_revision_documents"
    / "verify_revision_documents.py"
)
SPEC = importlib.util.spec_from_file_location("verify_revision_documents", MODULE_PATH)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


def build_supplement(
    path: Path,
    include_k6: bool = False,
    eicu_k3_selections: int = 1,
    figure_path: Path | None = None,
) -> None:
    document = Document()
    if figure_path is not None:
        paragraph = document.add_paragraph()
        paragraph.add_run().add_picture(str(figure_path))
    document.add_paragraph("Figure S2. Traceable candidate-K evidence.")
    table = document.add_table(rows=1, cols=6)
    for index, value in enumerate(
        [
            "Dataset",
            "Clusters (K)",
            "Seeds",
            "Mean deviance",
            "Mean high |lag-1| fraction",
            "Selections / 3 seeds",
        ]
    ):
        table.rows[0].cells[index].text = value
    for cohort in ["MIMIC-IV", "eICU-CRD", "AmsterdamUMCdb"]:
        for k in [2, 3, 4, 6 if include_k6 else 5]:
            cells = table.add_row().cells
            cells[0].text = cohort
            cells[1].text = str(k)
            cells[2].text = "3"
            cells[3].text = "100.000"
            cells[4].text = "0.000"
            cells[5].text = str(eicu_k3_selections if cohort == "eICU-CRD" and k == 3 else (2 if k == 3 else 0))
    document.save(path)


def write_fresh_k_summary(
    path: Path, include_k6: bool = False, eicu_k3_selections: int = 1
) -> None:
    rows = []
    cohort_map = {
        "MIMIC-IV": "mimic",
        "eICU-CRD": "eicu",
        "AmsterdamUMCdb": "aumc",
    }
    for cohort_label, cohort in cohort_map.items():
        for k in [2, 3, 4, 6 if include_k6 else 5]:
            rows.append(
                {
                    "cohort": cohort,
                    "K": k,
                    "seeds_completed": 3,
                    "mean_deviance": 100.0,
                    "mean_high_absolute_lag1_fraction": 0.0,
                    "selections_across_three_seeds": (
                        eicu_k3_selections
                        if cohort_label == "eICU-CRD" and k == 3
                        else (2 if k == 3 else 0)
                    ),
                }
            )
    pd.DataFrame(rows).to_csv(path, index=False)


def test_supplement_accepts_exact_traceable_k_grid(tmp_path):
    path = tmp_path / "supplement.docx"
    build_supplement(path)
    result = MODULE.supplement_checks(path)
    assert result["figure_s2_caption_count"] == 1
    assert result["table_s1_has_only_k2_k5"] is True
    assert result["table_s1_has_three_cohorts"] is True
    assert result["table_s1_selection_annotation_present"] is True
    assert result["table_s1_eicu_k3_not_unanimous"] is True


def test_supplement_rejects_k6(tmp_path):
    path = tmp_path / "supplement.docx"
    build_supplement(path, include_k6=True)
    assert MODULE.supplement_checks(path)["table_s1_has_only_k2_k5"] is False


def test_supplement_rejects_unanimous_eicu_k3_annotation(tmp_path):
    path = tmp_path / "supplement.docx"
    build_supplement(path, eicu_k3_selections=3)
    result = MODULE.supplement_checks(path)
    assert result["table_s1_eicu_k3_not_unanimous"] is False


def test_table_s1_matches_source_csv_cell_for_cell(tmp_path):
    supplement = tmp_path / "supplement.docx"
    source = tmp_path / "fresh_k_grid_summary.csv"
    build_supplement(supplement)
    write_fresh_k_summary(source)
    result = MODULE.supplement_checks(supplement, fresh_k_summary=source)
    identity = result["table_s1_source_identity"]
    assert identity["expected_cells"] == 78
    assert identity["mismatch_count"] == 0
    assert identity["exact_match"] is True


def test_table_s1_source_comparison_fails_on_one_changed_cell(tmp_path):
    supplement = tmp_path / "supplement.docx"
    source = tmp_path / "fresh_k_grid_summary.csv"
    build_supplement(supplement)
    write_fresh_k_summary(source)
    frame = pd.read_csv(source)
    frame.loc[0, "mean_deviance"] = 101.0
    frame.to_csv(source, index=False)
    identity = MODULE.supplement_checks(
        supplement, fresh_k_summary=source
    )["table_s1_source_identity"]
    assert identity["mismatch_count"] == 1
    assert identity["exact_match"] is False


def test_figure_s2_asset_identity_uses_sha256_when_bytes_match(tmp_path):
    figure = tmp_path / "figure.png"
    Image.new("RGB", (40, 30), (20, 40, 80)).save(figure)
    supplement = tmp_path / "supplement.docx"
    build_supplement(supplement, figure_path=figure)
    identity = MODULE.supplement_checks(
        supplement, figure_s2_source=figure
    )["figure_s2_asset_identity"]
    assert identity["identity_method"] == "sha256"
    assert identity["exact_match"] is True


def test_figure_s2_asset_identity_falls_back_to_pixel_comparison(tmp_path):
    source = tmp_path / "source.png"
    embedded = tmp_path / "embedded.png"
    image = Image.new("RGB", (40, 30), (20, 40, 80))
    image.save(source, compress_level=0)
    image.save(embedded, compress_level=9)
    assert source.read_bytes() != embedded.read_bytes()
    supplement = tmp_path / "supplement.docx"
    build_supplement(supplement, figure_path=embedded)
    identity = MODULE.supplement_checks(
        supplement, figure_s2_source=source
    )["figure_s2_asset_identity"]
    assert identity["identity_method"] == "pixel_comparison"
    assert identity["pixels_match"] is True
    assert identity["exact_match"] is True


def test_required_content_checks_cover_all_submission_guardrails():
    text = """
    AmsterdamUMCdb used the other three responses because BUN was unavailable.
    eICU-CRD retained initialization sensitivity.
    The ≥50%-coverage sensitivity failed because the PW component collapsed.
    The classifier showed no external incremental advantage over logistic regression.
    The diuretic analysis remains secondary and exploratory.
    Complete 30-window follow-up was cautionary, with material sensitivity to the
    missing-data rule. The screening grid did not select k=3 unanimously.
    """
    assert all(MODULE.required_content_checks(text).values())


def test_required_content_checks_fail_when_a_guardrail_is_missing():
    checks = MODULE.required_content_checks(
        "AmsterdamUMCdb used the other three responses; eICU retained initialization sensitivity."
    )
    assert checks["aumc_uses_three_clustering_variables"] is True
    assert checks["high_coverage_uo_does_not_preserve_pw"] is False
