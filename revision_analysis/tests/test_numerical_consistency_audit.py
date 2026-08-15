from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pandas as pd
from docx import Document


MODULE_PATH = (
    Path(__file__).resolve().parents[1]
    / "09_revision_documents"
    / "audit_numerical_consistency.py"
)
SPEC = importlib.util.spec_from_file_location("audit_numerical_consistency", MODULE_PATH)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


def test_anchor_tokens_are_required_in_the_same_anchored_paragraph():
    document = Document()
    document.add_paragraph("Primary result: n=123, OR 1.25 (95% CI 1.01–1.55).")
    checks = []
    MODULE.check_anchor_tokens(
        checks,
        check_id="result",
        category="test",
        source="source.csv",
        document_name="document.docx",
        document=document,
        anchor="Primary result",
        tokens=["n=123", "1.25", "1.01–1.55"],
        location="Results",
    )
    assert checks[0].passed is True

    checks = []
    MODULE.check_anchor_tokens(
        checks,
        check_id="result",
        category="test",
        source="source.csv",
        document_name="document.docx",
        document=document,
        anchor="Primary result",
        tokens=["n=123", "0.99"],
        location="Results",
    )
    assert checks[0].passed is False
    assert "0.99" in checks[0].details


def test_repeated_table_headers_select_the_requested_occurrence():
    document = Document()
    for value in ["first", "second"]:
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Feature"
        table.rows[0].cells[1].text = "Value"
        cells = table.add_row().cells
        cells[0].text = value
        cells[1].text = "1.000"

    frame = pd.DataFrame([{"Feature": "second", "Value": 1.0}])
    checks = []
    MODULE.check_dataframe_table(
        checks,
        check_id="second_table",
        category="test",
        source="source.csv",
        document_name="document.docx",
        document=document,
        location="second table",
        frame=frame,
        occurrence=1,
    )
    assert checks[0].passed is True
    assert '"matching_tables": 2' in checks[0].details
    assert '"selected_occurrence": 1' in checks[0].details


def test_table_identity_fails_closed_when_occurrence_is_missing():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Feature"
    table.rows[0].cells[1].text = "Value"
    frame = pd.DataFrame([{"Feature": "row", "Value": 1.0}])
    checks = []
    MODULE.check_dataframe_table(
        checks,
        check_id="missing_table",
        category="test",
        source="source.csv",
        document_name="document.docx",
        document=document,
        location="missing table",
        frame=frame,
        occurrence=1,
    )
    assert checks[0].passed is False


def test_grid_mismatch_reports_exact_cell_coordinates():
    mismatches = MODULE.grid_mismatches(
        [["A", "B"], ["x", "1.000"]],
        [["A", "B"], ["x", "1.001"]],
    )
    assert mismatches == [
        {"row": 1, "column": 1, "expected": "1.000", "observed": "1.001"}
    ]


def test_table_s2_eicu_audit_fails_on_a_mixed_denominator_cell(tmp_path):
    source = tmp_path / "table_s2.csv"
    pd.DataFrame(
        [
            {
                "Characteristic": "Patient number, N",
                "Overall": "1417",
                "RR": "869",
                "DR": "423",
                "PW": "125",
                "P value": "",
            },
            {
                "Characteristic": "Myocardial infarct, n (%)",
                "Overall": "47 (3.3)",
                "RR": "20 (2.3)",
                "DR": "19 (4.5)",
                "PW": "8 (6.4)",
                "P value": "0.016",
            },
        ]
    ).to_csv(source, index=False)
    document = Document()
    table = document.add_table(rows=2, cols=12)
    for label, values in [
        ("Patient number, N", ["1417", "869", "423", "125", ""]),
        ("Myocardial infarct, n (%)", ["62 (3.1)", "20 (2.3)", "19 (4.5)", "8 (6.4)", "0.016"]),
    ]:
        cells = table.add_row().cells
        cells[0].text = label
        for cell, value in zip(cells[7:12], values):
            cell.text = value
    checks = []
    MODULE.check_table_s2_eicu(checks, document=document, source_path=source)
    assert checks[0].passed is False
    assert "62 (3.1)" in checks[0].details

    table.rows[3].cells[7].text = "47 (3.3)"
    checks = []
    MODULE.check_table_s2_eicu(checks, document=document, source_path=source)
    assert checks[0].passed is True
