from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pandas as pd
import pytest
from docx import Document


MODULE_PATH = (
    Path(__file__).resolve().parents[1]
    / "09_revision_documents"
    / "repair_reporting_consistency.py"
)
SPEC = importlib.util.spec_from_file_location("repair_reporting_consistency", MODULE_PATH)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


def count_cell(count: int, denominator: int) -> str:
    return f"{count} ({100 * count / denominator:.1f})"


def write_characteristics(path: Path, *, overall_n: int = 1417) -> None:
    denominators = {"Overall": overall_n, "C1": 423, "C2": 869, "C3": 125}
    rows = [
        {
            "feature": "n",
            "type": "",
            "Overall": str(overall_n),
            "C1": "423",
            "C2": "869",
            "C3": "125",
            "P-Value": "",
        }
    ]
    features = [
        "Myocardial infarct, n (%)",
        "Congestive heart failure, n (%)",
        "Hepatic, n (%)",
        "Diabetes, n (%)",
        "Hypertension, n (%)",
        "Respiratory failure, n (%)",
        "Septic shock, n (%)",
        "Acidosis, n (%)",
        "pulmonary, n (%)",
    ]
    for index, feature in enumerate(features, start=1):
        counts = {"Overall": index, "C1": index, "C2": index, "C3": index}
        rows.append(
            {
                "feature": feature,
                "type": "1",
                **{
                    column: count_cell(counts[column], denominators[column])
                    for column in denominators
                },
                "P-Value": "0.500",
            }
        )
    pd.DataFrame(rows).to_csv(path, index=False)


def test_table_s2_uses_current_denominators_and_never_substitutes_broad_pulmonary(
    tmp_path: Path,
) -> None:
    source = tmp_path / "characteristics.csv"
    write_characteristics(source)
    result = MODULE.authoritative_table_s2(source)
    chronic = result.loc[
        result.Characteristic.eq("Chronic pulmonary disease, n (%)")
    ].iloc[0]
    assert chronic["Overall"] == "—"
    assert chronic["RR"] == "—"
    assert chronic["Source field"] == "Not available"
    assert "pulmonary, n (%)" not in set(result["Source field"])
    denominator_row = result.iloc[0]
    assert denominator_row[["Overall", "RR", "DR", "PW"]].tolist() == [
        "1417",
        "869",
        "423",
        "125",
    ]


def test_table_s2_rejects_legacy_overall_denominator(tmp_path: Path) -> None:
    source = tmp_path / "legacy.csv"
    write_characteristics(source, overall_n=1970)
    with pytest.raises(ValueError, match="not the authoritative N=1,417 cohort"):
        MODULE.authoritative_table_s2(source)


def test_table_s2_rejects_percentage_denominator_mismatch(tmp_path: Path) -> None:
    source = tmp_path / "bad_percentage.csv"
    write_characteristics(source)
    frame = pd.read_csv(source, dtype=str, keep_default_na=False)
    frame.loc[frame.feature.eq("Diabetes, n (%)"), "C1"] = "10 (99.9)"
    frame.to_csv(source, index=False)
    with pytest.raises(ValueError, match="Percentage/denominator mismatch"):
        MODULE.authoritative_table_s2(source)


def test_table_s19_contains_all_three_definitions_for_each_cohort(tmp_path: Path) -> None:
    rows = []
    for cohort in ("MIMIC-IV", "AUMC"):
        for index, definition in enumerate(MODULE.RESPONSE_LABELS, start=1):
            rows.append(
                {
                    "cohort_label": cohort,
                    "response_definition": definition,
                    "adjusted_or_response_vs_nonresponse": 0.5 + index / 10,
                    "ci_low": 0.3,
                    "ci_high": 1.1,
                    "p_value": 0.05,
                    "n_complete": 100,
                    "deaths": 10,
                    "responsive_n": 60,
                }
            )
    source = tmp_path / "diuretic.csv"
    pd.DataFrame(rows).to_csv(source, index=False)
    result = MODULE.diuretic_table_s19(source)
    assert len(result) == 6
    assert set(result.Definition) == set(MODULE.RESPONSE_LABELS.values())

    document = Document()
    table = document.add_table(rows=1, cols=len(result.columns))
    for cell, value in zip(table.rows[0].cells, result.columns):
        cell.text = value
    for row in result.loc[~result.Definition.eq("Absolute ≥200 mL")].itertuples(
        index=False, name=None
    ):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
    document.add_paragraph("Figure S14. Old caption.")
    MODULE.patch_table_s19(document, result)
    assert len(table.rows) == 7
    assert [row.cells[1].text for row in table.rows[1:]].count("Absolute ≥200 mL") == 2
    assert "three" not in document.paragraphs[-1].text.lower()
    assert "absolute ≥200 mL" in document.paragraphs[-1].text
